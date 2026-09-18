"""Focused tests for CV Generation Quality v2 Task 4 DOCX rendering."""

from __future__ import annotations

import copy
import hashlib
import time
import unittest
import zipfile
from io import BytesIO

from docx import Document

from product.cv_document_model import build_cv_document_model
from product.cv_document_renderer import (
    CV_DOCUMENT_RENDERER_VERSION,
    CvDocumentRendererError,
    _freeze_docx_bytes,
    render_cv_document,
)


def statement(
    statement_id: str,
    target_section: str,
    text: str,
    *,
    record_id: str | None = None,
    evidence_ids: list[str] | None = None,
    requirement_ids: list[str] | None = None,
) -> dict:
    evidence_ids = evidence_ids or [f"clm_{statement_id}"]
    requirement_ids = requirement_ids or []
    return {
        "statement_id": statement_id,
        "target_section": target_section,
        "target_record_id": record_id,
        "source_profile_evidence_ids": evidence_ids,
        "source_job_requirement_ids": requirement_ids,
        "statement_text": text,
        "structured_content": {
            "fragments": [
                {"profile_evidence_id": evidence_id, "record_id": record_id, "text": text}
                for evidence_id in evidence_ids
            ]
        },
        "provenance": {
            "profile_evidence_ids": evidence_ids,
            "job_requirement_ids": requirement_ids,
        },
    }


def statement_plan(statements: list[dict], *, gaps: list[dict] | None = None) -> dict:
    return {
        "schema_version": "cv-statement-plan.v0",
        "statements": copy.deepcopy(statements),
        "omitted_evidence": [],
        "ungroupable_evidence": [],
        "uncovered_requirements": gaps or [],
        "provenance": {
            "statement_profile_evidence_ids": sorted(
                {
                    evidence_id
                    for item in statements
                    for evidence_id in item.get("source_profile_evidence_ids", [])
                }
            ),
            "source_selected_job_requirement_ids": sorted(
                {
                    requirement_id
                    for item in statements
                    for requirement_id in item.get("source_job_requirement_ids", [])
                }
            ),
        },
    }


def rich_document_model() -> dict:
    plan = statement_plan(
        [
            statement(
                "stmt_summary",
                "professional_summary",
                "Delivered Python ETL pipelines across UK/North Sea assets — reducing NPT by 15%.",
            ),
            statement("stmt_skill_1", "skills", "Python"),
            statement("stmt_skill_2", "skills", "SQL"),
            statement(
                "stmt_role_a_1",
                "professional_experience",
                "Built ETL pipelines for drilling operations.",
                record_id="rec_a",
            ),
            statement(
                "stmt_role_a_2",
                "professional_experience",
                "Managed R&D operations across UK/North Sea assets.",
                record_id="rec_a",
            ),
            statement(
                "stmt_role_b_1",
                "professional_experience",
                "Delivered on-call incident response for production systems.",
                record_id="rec_b",
            ),
            statement("stmt_qual_1", "qualifications", "MSc Engineering"),
        ],
        gaps=[
            {
                "job_requirement_id": "job_german",
                "text": "Requires German language proficiency",
                "kind": "required",
                "source": "resolved_job_evidence",
                "gap": None,
            }
        ],
    )
    return build_cv_document_model(plan)


def sparse_document_model() -> dict:
    plan = statement_plan([statement("stmt_skill", "skills", "Python")])
    return build_cv_document_model(plan)


def missing_optional_employment_fields_model() -> dict:
    # This document model schema has no employer/title/date/location fields
    # at all -- a role is identified only by record_id. This fixture proves
    # the renderer displays the role's statements without fabricating any
    # employer/title/date/location text the model never carried.
    plan = statement_plan(
        [
            statement(
                "stmt_role_only",
                "professional_experience",
                "Delivered infrastructure automation.",
                record_id="rec_only",
            ),
        ]
    )
    return build_cv_document_model(plan)


def text_fidelity_model() -> dict:
    text = "Managed R&D operations across UK/North Sea assets — reducing NPT by 15%."
    plan = statement_plan([statement("stmt_fidelity", "professional_summary", text)])
    return build_cv_document_model(plan), text


def visible_paragraphs(docx_bytes: bytes) -> list[str]:
    document = Document(BytesIO(docx_bytes))
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    for section in document.sections:
        for part in (section.header, section.footer):
            for paragraph in part.paragraphs:
                paragraphs.append(paragraph.text)
    return paragraphs


def all_visible_text(docx_bytes: bytes) -> str:
    document = Document(BytesIO(docx_bytes))
    chunks = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                chunks.append(cell.text)
    for section in document.sections:
        for part in (section.header, section.footer):
            for paragraph in part.paragraphs:
                chunks.append(paragraph.text)
    return "\n".join(chunks)


class TestCvDocumentRendererCore(unittest.TestCase):
    def test_valid_model_produces_a_readable_docx(self):
        docx_bytes = render_cv_document(rich_document_model())
        document = Document(BytesIO(docx_bytes))
        self.assertTrue(len(document.paragraphs) > 0)

    def test_sections_appear_in_deterministic_model_order(self):
        model = rich_document_model()
        docx_bytes = render_cv_document(model)
        text = all_visible_text(docx_bytes)
        expected_titles = [section["title"] for section in model["sections"]]
        positions = [text.index(title) for title in expected_titles]
        self.assertEqual(positions, sorted(positions))

    def test_evidence_gaps_section_is_rendered_visibly(self):
        model = rich_document_model()
        docx_bytes = render_cv_document(model)
        text = all_visible_text(docx_bytes)
        self.assertIn("Evidence Gaps", text)
        self.assertIn("Requires German language proficiency", text)

    def test_statement_text_is_rendered_verbatim(self):
        model = rich_document_model()
        docx_bytes = render_cv_document(model)
        text = all_visible_text(docx_bytes)
        self.assertIn(
            "Delivered Python ETL pipelines across UK/North Sea assets — reducing NPT by 15%.",
            text,
        )
        self.assertIn("Managed R&D operations across UK/North Sea assets.", text)
        self.assertIn("MSc Engineering", text)

    def test_employment_roles_preserve_model_order(self):
        model = rich_document_model()
        docx_bytes = render_cv_document(model)
        text = all_visible_text(docx_bytes)
        first_statement = "Built ETL pipelines for drilling operations."
        second_statement = "Delivered on-call incident response for production systems."
        self.assertLess(text.index(first_statement), text.index(second_statement))

    def test_bullets_within_a_role_preserve_statement_order(self):
        model = rich_document_model()
        docx_bytes = render_cv_document(model)
        text = all_visible_text(docx_bytes)
        first = "Built ETL pipelines for drilling operations."
        second = "Managed R&D operations across UK/North Sea assets."
        self.assertLess(text.index(first), text.index(second))

    def test_text_fidelity_unicode_and_punctuation_survive(self):
        model, expected_text = text_fidelity_model()
        docx_bytes = render_cv_document(model)
        text = all_visible_text(docx_bytes)
        self.assertIn(expected_text, text)


class TestCvDocumentRendererOmission(unittest.TestCase):
    def test_optional_empty_sections_disappear_cleanly(self):
        docx_bytes = render_cv_document(sparse_document_model())
        text = all_visible_text(docx_bytes)
        for absent_heading in ("Professional Summary", "Professional Experience", "Qualifications"):
            self.assertNotIn(absent_heading, text)

    def test_sparse_model_still_renders_present_section(self):
        docx_bytes = render_cv_document(sparse_document_model())
        text = all_visible_text(docx_bytes)
        self.assertIn("Python", text)

    def test_missing_optional_employment_fields_are_not_invented(self):
        docx_bytes = render_cv_document(missing_optional_employment_fields_model())
        text = all_visible_text(docx_bytes)
        self.assertIn("Delivered infrastructure automation.", text)
        for invented in ("Present", "Current", "Employer", "Company"):
            self.assertNotIn(invented, text)

    def test_no_placeholder_contact_or_identity_text_is_added(self):
        docx_bytes = render_cv_document(rich_document_model())
        text = all_visible_text(docx_bytes)
        for placeholder in ("Candidate Name", "Your Name", "Name unavailable", "Email unavailable"):
            self.assertNotIn(placeholder, text)


class TestCvDocumentRendererProvenanceLeakage(unittest.TestCase):
    def test_internal_evidence_and_statement_ids_do_not_appear_visibly(self):
        model = rich_document_model()
        docx_bytes = render_cv_document(model)
        text = all_visible_text(docx_bytes)
        for section in model["sections"]:
            if section["section_id"] == "professional_experience":
                for role in section["roles"]:
                    self.assertNotIn(role["record_id"], text)
                    self.assertNotIn(role["role_id"], text)
                    for item in role["statements"]:
                        self.assertNotIn(item["statement_id"], text)
                        for evidence_id in item["source_profile_evidence_ids"]:
                            self.assertNotIn(evidence_id, text)
            elif "items" in section:
                for item in section["items"]:
                    if item.get("statement_id"):
                        self.assertNotIn(item["statement_id"], text)
                    for evidence_id in item.get("source_profile_evidence_ids", []):
                        self.assertNotIn(evidence_id, text)

    def test_diagnostics_and_internal_metadata_do_not_appear_visibly(self):
        docx_bytes = render_cv_document(rich_document_model())
        text = all_visible_text(docx_bytes).lower()
        for leaked_term in (
            "cv-document-model.v2",
            "provenance",
            "structural_decision",
            "budget",
            "suppressed",
            "decision_metadata",
        ):
            self.assertNotIn(leaked_term, text)

    def test_unmistakable_planted_ids_never_appear(self):
        model = rich_document_model()
        planted_claim_id = "clm_renderer_leak_test_001"
        planted_statement_id = "stmt_renderer_leak_test_001"
        planted_record_id = "rec_renderer_leak_test_001"

        role_section = next(s for s in model["sections"] if s["section_id"] == "professional_experience")
        role_section["roles"][0]["record_id"] = planted_record_id
        role_section["roles"][0]["role_id"] = planted_record_id
        role_section["roles"][0]["statements"][0]["statement_id"] = planted_statement_id
        role_section["roles"][0]["statements"][0]["target_record_id"] = planted_record_id
        role_section["roles"][0]["statements"][0]["source_profile_evidence_ids"] = [planted_claim_id]

        docx_bytes = render_cv_document(model)
        text = all_visible_text(docx_bytes)
        self.assertNotIn(planted_claim_id, text)
        self.assertNotIn(planted_statement_id, text)
        self.assertNotIn(planted_record_id, text)


class TestCvDocumentRendererImmutability(unittest.TestCase):
    def test_rendering_does_not_mutate_the_input_model(self):
        model = rich_document_model()
        before = copy.deepcopy(model)
        render_cv_document(model)
        self.assertEqual(model, before)


class TestCvDocumentRendererDeterminism(unittest.TestCase):
    def test_identical_input_renders_identical_bytes_immediately(self):
        model = rich_document_model()
        first = render_cv_document(model)
        second = render_cv_document(copy.deepcopy(model))
        self.assertEqual(first, second)

    def test_identical_input_renders_identical_bytes_across_wall_clock_gap(self):
        model = rich_document_model()
        first = render_cv_document(model)
        time.sleep(1.2)
        second = render_cv_document(model)
        self.assertEqual(first, second)
        self.assertEqual(hashlib.sha256(first).digest(), hashlib.sha256(second).digest())
        Document(BytesIO(first))
        Document(BytesIO(second))

    def test_meaningful_content_change_changes_bytes_and_hash(self):
        model = rich_document_model()
        first = render_cv_document(model)

        changed = copy.deepcopy(model)
        summary_section = next(s for s in changed["sections"] if s["section_id"] == "professional_summary")
        summary_section["items"][0]["statement_text"] = "A materially different summary statement."
        second = render_cv_document(changed)

        self.assertNotEqual(first, second)
        self.assertNotEqual(hashlib.sha256(first).digest(), hashlib.sha256(second).digest())


class TestCvDocumentRendererZipFreezing(unittest.TestCase):
    def test_zip_freezing_is_host_platform_independent(self):
        base_bytes = render_cv_document(rich_document_model())
        source = zipfile.ZipFile(BytesIO(base_bytes))

        def archive_with_creator_system(create_system: int) -> bytes:
            buffer = BytesIO()
            with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
                for item in source.infolist():
                    info = zipfile.ZipInfo(item.filename, date_time=(2026, 1, 1, 0, 0, 0))
                    info.compress_type = item.compress_type
                    info.external_attr = item.external_attr
                    info.create_system = create_system
                    archive.writestr(info, source.read(item.filename))
            return buffer.getvalue()

        windows_archive = archive_with_creator_system(0)
        unix_archive = archive_with_creator_system(3)

        frozen_windows = _freeze_docx_bytes(windows_archive)
        frozen_unix = _freeze_docx_bytes(unix_archive)

        self.assertEqual(frozen_windows, frozen_unix)


class TestCvDocumentRendererSchemaGate(unittest.TestCase):
    def test_unknown_schema_version_is_rejected(self):
        model = rich_document_model()
        model["schema_version"] = "cv-document-model.v99"
        with self.assertRaises(CvDocumentRendererError):
            render_cv_document(model)

    def test_missing_schema_version_is_rejected(self):
        model = rich_document_model()
        del model["schema_version"]
        with self.assertRaises(CvDocumentRendererError):
            render_cv_document(model)

    def test_malformed_model_is_rejected_behind_a_stable_error(self):
        with self.assertRaises(CvDocumentRendererError):
            render_cv_document({"schema_version": "cv-document-model.v2", "sections": "not-a-list"})

    def test_non_dict_input_is_rejected(self):
        with self.assertRaises(CvDocumentRendererError):
            render_cv_document(["not", "a", "model"])


class TestCvDocumentRendererImportBoundary(unittest.TestCase):
    def test_renderer_module_has_no_forbidden_upstream_dependencies(self):
        import ast
        import pathlib

        module_path = pathlib.Path(__file__).resolve().parents[1] / "product" / "cv_document_renderer.py"
        tree = ast.parse(module_path.read_text(encoding="utf-8"))

        imported_modules: list[str] = []
        for node in ast.walk(tree):
            if isinstance(node, ast.Import):
                imported_modules.extend(alias.name for alias in node.names)
            elif isinstance(node, ast.ImportFrom) and node.module:
                imported_modules.append(node.module)

        forbidden_substrings = (
            "webapp",
            "application_pack_renderer",
            "application_pack_contract",
            "application_pack",
            "profile",
            "job_fit",
            "application_intelligence",
            "llm",
            "provider",
            "database",
            "persistence",
            "sqlalchemy",
            "cv_content_plan",
            "cv_statement_plan",
        )
        for module_name in imported_modules:
            lowered = module_name.lower()
            for forbidden in forbidden_substrings:
                self.assertNotIn(
                    forbidden,
                    lowered,
                    msg=f"renderer must not import {module_name!r} (matched {forbidden!r})",
                )

    def test_renderer_does_not_import_application_pack_renderer_module(self):
        import product.cv_document_renderer as renderer_module

        self.assertNotIn("product.application_pack_renderer", vars(renderer_module))
        self.assertNotIn("application_pack_renderer", renderer_module.__dict__)


class TestCvDocumentRendererVersion(unittest.TestCase):
    def test_renderer_version_constant_is_exposed(self):
        self.assertTrue(CV_DOCUMENT_RENDERER_VERSION)


if __name__ == "__main__":
    unittest.main()
