"""Real Chromium + Uvicorn browser journeys for the Ticket 9 product UI."""
from __future__ import annotations

import json
import socket
import threading
import time
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
from urllib.parse import urlsplit

import pytest
import uvicorn
from docx import Document

from product.application_document_contract import DOCX_MEDIA_TYPE
from product.application_intelligence_providers import ProviderResponse as AIResponse
from product.job_understanding_providers import ProviderResponse as UnderstandingResponse
from webapp.app import create_app
from webapp.config import Settings
from webapp.persistence.accounts import create_account
from webapp.persistence.application_documents import create_document_version, get_selection
from webapp.persistence.artifacts import get_current_artifact, list_artifact_history
from webapp.persistence.db import connect
from webapp.persistence.search_workspaces import create_search_workspace
from webapp.persistence.workspaces import PROFILE_WORKSPACE_ID, create_workspace
from webapp.services.document_blob_store import DocumentBlobStore

from tests.webapp.fixtures.acceptance.fixtures import extension
from webapp.persistence.accounts import DEFAULT_ACCOUNT_ID


POSTING_TEXT = (
    "Python is required.\n"
    "Cloud certification is required.\n"
    "Build reliable data pipelines.\n"
    "Applicants must already have the right to work in the UK.\n"
    "German would be an advantage.\n"
    "Hybrid role: two days per week in London.\n"
)


def _free_port() -> int:
    with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as sock:
        sock.bind(("127.0.0.1", 0))
        return sock.getsockname()[1]


def _write_profile_root(root: Path) -> None:
    candidate = root / ".claude/skills/job-application-assistant"
    candidate.mkdir(parents=True)
    (root / "cv").mkdir(parents=True)
    (root / "CLAUDE.md").write_text(
        """# Job Application Assistant for Ada Lovelace

## Candidate Profile

### Identity
- **Name:** Ada Lovelace
- **Location:** London hybrid
- **Languages:**

| Language | Level |
|----------|-------|
| German | Professional |

- **Status:** Employed

### Professional Experience
- **Data Engineer** (2020-01 - Present) - **Evidence Works** (London)

### Technical Skills
- **Primary:** Python
""",
        encoding="utf-8",
    )
    (candidate / "01-candidate-profile.md").write_text(
        """# Candidate Profile

## Identity
- **Name:** Ada Lovelace
- **Location:** London hybrid
- **Status:** Employed
- **Constraints:** Right to work in the UK

### Languages

| Language | Level | Notes |
|----------|-------|-------|
| German | Professional | professional use |

## Professional Experience

### Data Engineer - Evidence Works (2020-01 - Present)
London
- Built production data pipelines
- Coordinated complex engineering schedules resources risks milestones recovery actions reporting delivery planning controls stakeholder communication governance assurance oversight
- I bring evidence backed project planning experience across complex engineering operations coordinating schedules resources risks milestones recovery actions field teams leadership reporting data analysis delivery governance quality controls stakeholder communication continuous improvement operational readiness tender planning lessons learned and critical path protection

## Technical Skills

### Programming & ML
- Python

## Publications
1. Ada Lovelace (2026). Notes on the Analytical Engine.
""",
        encoding="utf-8",
    )
    (root / "cv/main_example.tex").write_text(
        "\\documentclass{moderncv}\\name{Ada}{Lovelace}\\begin{document}\\end{document}\n",
        encoding="utf-8",
    )


class _UnderstandingProvider:
    provider_id = "browser-fake"
    model_id = "browser-fixture"
    model_version = "v0"

    def extract(self, request):
        quotes = (
            ("python", "requirements", "required", "Python is required."),
            ("cloud", "requirements", "required", "Cloud certification is required."),
            ("pipelines", "responsibilities", "required", "Build reliable data pipelines."),
            ("rights", "eligibility_requirements", "required", "Applicants must already have the right to work in the UK."),
            ("german", "language_requirements", "preferred", "German would be an advantage."),
            ("hybrid", "logistics_requirements", "required", "Hybrid role: two days per week in London."),
        )
        return UnderstandingResponse(payload={
            "schema_version": "job-understanding-candidate.v0",
            "items": [
                {
                    "proposal_id": f"browser-{name}", "category": category,
                    "kind": kind, "quote": quote, "certainty": "explicit",
                }
                for name, category, kind, quote in quotes
            ],
            "suggestions": [], "ambiguous_statements": [], "warnings": [],
        })


def _claim_id(claims: list[dict], needle: str) -> str:
    return next(
        claim["id"] for claim in claims
        if needle.casefold() in str(claim.get("value", "")).casefold()
    )


def _job_id(evidence: list[dict], exact_text: str) -> str:
    return next(item["id"] for item in evidence if item["text"] == exact_text)


class _SemanticAdapter:
    """Dynamic fake: references the real IDs produced by Profile/Ticket 6."""

    def propose(self, *, profile_evidence, resolved_job_evidence, active_extensions):
        evidence = resolved_job_evidence["evidence"]
        python_claim = _claim_id(profile_evidence, "Python")
        pipeline_claim = _claim_id(profile_evidence, "Built production data pipelines")
        rights_claim = _claim_id(profile_evidence, "Right to work")
        german_claim = _claim_id(profile_evidence, "German")
        location_claim = _claim_id(profile_evidence, "London hybrid")
        return {
            "matches": [
                {
                    "proposal_id": "browser-direct",
                    "job_evidence_id": _job_id(evidence, "Python is required."),
                    "profile_evidence_ids": [python_claim],
                    "classification": "direct",
                    "rationale": "Explicit Python evidence on both sides.",
                    "confidence": "high",
                },
                {
                    "proposal_id": "browser-functional",
                    "job_evidence_id": _job_id(evidence, "Build reliable data pipelines."),
                    "profile_evidence_ids": [pipeline_claim],
                    "classification": "functionally_equivalent",
                    "rationale": "The responsibilities align by function.",
                    "confidence": "high",
                    "functional_basis": {
                        "responsibility_alignment": [
                            "Build reliable data pipelines",
                            "Built production data pipelines",
                        ],
                        "competency_alignment": [],
                        "title_similarity_only": False,
                    },
                },
                {
                    "proposal_id": "browser-transfer",
                    "job_evidence_id": _job_id(evidence, "German would be an advantage."),
                    "profile_evidence_ids": [pipeline_claim],
                    "classification": "transferable",
                    "rationale": "A bounded active mapping was proposed.",
                    "confidence": "medium",
                    "extension_ref": {
                        "extension_id": "data-transfer",
                        "extension_version": "0.1.0",
                        "record_type": "transferable_mapping",
                        "record_id": "field-models-to-pipelines",
                    },
                },
            ],
            "gates": [
                {
                    "gate_id": "eligibility", "status": "PASS",
                    "reason": "Affirmative candidate and job evidence.",
                    "job_evidence_ids": [_job_id(evidence, "Applicants must already have the right to work in the UK.")],
                    "profile_evidence_ids": [rights_claim],
                },
                {
                    "gate_id": "language", "status": "PASS",
                    "reason": "Affirmative candidate and job evidence.",
                    "job_evidence_ids": [_job_id(evidence, "German would be an advantage.")],
                    "profile_evidence_ids": [german_claim],
                },
                {
                    "gate_id": "location_logistics", "status": "PASS",
                    "reason": "Affirmative candidate and job evidence.",
                    "job_evidence_ids": [_job_id(evidence, "Hybrid role: two days per week in London.")],
                    "profile_evidence_ids": [location_claim],
                },
            ],
        }


class _ApplicationIntelligenceProvider:
    provider_id = "browser-fake"
    model_id = "browser-fixture"
    model_version = "v0"

    def propose(self, request):
        python_claim = _claim_id(request["profile_snapshot"]["claims"], "Python")
        summary_claim = _claim_id(
            request["profile_snapshot"]["claims"], "Coordinated complex engineering schedules"
        )
        cover_claim = _claim_id(
            request["profile_snapshot"]["claims"], "I bring evidence backed project planning"
        )
        valid_atom = {
            "atom_id": "browser-valid", "atom_kind": "candidate_fact",
            "assertion_type": "technical_skill",
            "profile_evidence_ids": [python_claim], "rendering_variant": "PLAIN",
        }
        unknown_atom = {
            "atom_id": "browser-unsupported", "atom_kind": "candidate_fact",
            "assertion_type": "certification",
            "profile_evidence_ids": ["clm_9999999999999999"],
            "rendering_variant": "PLAIN",
        }
        return AIResponse(payload={"content_units": [
            {
                "unit_id": "cv-ready", "unit_type": "cv_bullet",
                "atoms": [valid_atom], "connectives": [],
            },
            {
                "unit_id": "cv-needs-review", "unit_type": "cv_bullet",
                "atoms": [valid_atom, unknown_atom], "connectives": [],
            },
            {
                "unit_id": "cv-summary-ready", "unit_type": "cv_summary_line",
                "atoms": [{
                    "atom_id": "browser-summary", "atom_kind": "candidate_fact",
                    "assertion_type": "responsibility",
                    "profile_evidence_ids": [summary_claim], "rendering_variant": "PLAIN",
                }],
                "connectives": [],
            },
            {
                "unit_id": "cover-ready", "unit_type": "cover_letter_paragraph",
                "atoms": [{
                    "atom_id": "browser-cover", "atom_kind": "candidate_fact",
                    "assertion_type": "responsibility",
                    "profile_evidence_ids": [cover_claim], "rendering_variant": "PLAIN",
                }],
                "connectives": [],
            },
            {
                "unit_id": "cv-unsupported-only", "unit_type": "cv_bullet",
                "atoms": [unknown_atom], "connectives": [],
            },
        ]})


class _DiscoveryRunner:
    def search(self, source, **kwargs):
        assert source == "freehire-search"
        return [{
            "id": "browser-discovery-1",
            "title": "Evidence Data Engineer",
            "company": "Discovery Evidence Co",
            "location": "London, UK",
            "date": "2026-08-20",
            "url": "https://freehire.me/jobs/browser-discovery-1",
            "work_mode": "hybrid",
            "regions": ["eu"],
            "countries": ["GB"],
            "skills": ["Python"],
            "description": POSTING_TEXT,
        }]


@pytest.fixture
def live_server(tmp_path, monkeypatch):
    profile_root = tmp_path / "profile"
    _write_profile_root(profile_root)
    extensions_dir = tmp_path / "private-extension-registry"
    extension_dir = extensions_dir / "data-transfer"
    extension_dir.mkdir(parents=True)
    (extension_dir / "extension.json").write_text(
        json.dumps(extension(conditional=True)), encoding="utf-8"
    )
    browser_secret = "sk-browser-must-never-render"
    monkeypatch.setenv("OPENAI_API_KEY", browser_secret)
    port = _free_port()
    settings = Settings(
        db_path=tmp_path / "browser.sqlite3", host="127.0.0.1", port=port,
        profile_root=str(profile_root), extensions_dir=extensions_dir,
        documents_root=tmp_path / "documents",
    )
    app = create_app(settings)
    app.state.job_understanding_provider = _UnderstandingProvider()
    app.state.semantic_adapter = _SemanticAdapter()
    app.state.application_intelligence_provider = _ApplicationIntelligenceProvider()
    app.state.discovery_portal_runner = _DiscoveryRunner()
    server = uvicorn.Server(uvicorn.Config(
        app, host="127.0.0.1", port=port, log_level="warning", access_log=False,
    ))
    thread = threading.Thread(target=server.run, daemon=True)
    thread.start()
    deadline = time.monotonic() + 15
    while time.monotonic() < deadline:
        try:
            with socket.create_connection(("127.0.0.1", port), timeout=0.25):
                break
        except OSError:
            time.sleep(0.05)
    else:
        server.should_exit = True
        thread.join(timeout=5)
        raise RuntimeError("Uvicorn browser fixture did not start on 127.0.0.1")
    yield SimpleNamespace(
        base_url=f"http://127.0.0.1:{port}", profile_root=profile_root,
        extensions_dir=extensions_dir, secret=browser_secret, db_path=settings.db_path,
    )
    server.should_exit = True
    thread.join(timeout=10)
    assert not thread.is_alive(), "Uvicorn browser fixture did not stop cleanly"


def _click_reload(page, locator) -> None:
    with page.expect_navigation(wait_until="networkidle"):
        locator.click()


def _refresh_profile(page, live_server) -> None:
    page.goto(f"{live_server.base_url}/profile", wait_until="networkidle")
    _click_reload(page, page.get_by_role("button", name="Refresh snapshot"))
    assert page.get_by_text("Verified evidence").first.is_visible()


def _create_job(page, live_server, company="Browser Evidence Co") -> str:
    page.goto(f"{live_server.base_url}/new-job", wait_until="networkidle")
    paste_panel = page.locator('[data-mode-panel="paste"]')
    manual_panel = page.locator('[data-mode-panel="manual"]')
    import_panel = page.locator('[data-mode-panel="import"]')
    assert paste_panel.is_visible()
    assert not manual_panel.is_visible()
    assert not import_panel.is_visible()
    page.locator('input[name="mode"][value="manual"]').check()
    assert not paste_panel.is_visible()
    assert manual_panel.is_visible()
    assert not import_panel.is_visible()
    page.locator('input[name="mode"][value="import"]').check()
    assert not paste_panel.is_visible()
    assert not manual_panel.is_visible()
    assert import_panel.is_visible()
    page.locator('input[name="mode"][value="paste"]').check()
    page.locator('input[name="company"]').fill(company)
    page.locator('input[name="title"]').fill("Evidence Data Engineer")
    page.locator('textarea[name="posting_text"]').fill(POSTING_TEXT)
    with page.expect_navigation(wait_until="networkidle"):
        page.get_by_role("button", name="Create workspace").click()
    return page.url


def _run_to_intelligence(page, live_server) -> str:
    workspace_url = _create_job(page, live_server)
    _click_reload(page, page.get_by_role("button", name="Run Understanding"))
    assert page.get_by_text("Accepted job evidence", exact=True).count() == 6
    page.locator('input[name="extension_ids"][value="data-transfer"]').check()
    _click_reload(page, page.get_by_role("button", name="Run Job Fit"))
    page.get_by_text("Technical details: evidence matches, gaps, and IDs", exact=True).click()
    assert page.get_by_text("Verified evidence", exact=True).is_visible()
    assert page.get_by_text("Accepted inference — functionally equivalent", exact=True).is_visible()
    assert page.get_by_text("Transferable evidence", exact=True).is_visible()
    assert page.get_by_text("Missing evidence", exact=True).is_visible()
    assert page.get_by_text("Functional basis:").is_visible()
    assert page.get_by_text("Candidate evidence exists").is_visible()
    assert page.get_by_text("Does not prove employment history").is_visible()
    assert page.get_by_text("NEEDS_REVIEW", exact=True).first.is_visible()
    _click_reload(page, page.get_by_role("button", name="Run Application Intelligence"))
    return workspace_url


def _resolve_all_pending_reviews(page, disposition: str) -> None:
    for _ in range(30):
        button = page.locator(
            'article.review-item:not(:has(.decision)) '
            f'button.review-action[data-disposition="{disposition}"]'
        ).first
        if button.count() == 0:
            break
        _click_reload(page, button)
    else:
        raise AssertionError(
            f"review queue did not converge for disposition={disposition!r}"
        )


def _confirm_pack(page) -> None:
    _click_reload(page, page.get_by_role("button", name="Generate AI documents"))
    for kind in ("cv", "cover_letter"):
        panel = page.locator(f'[data-document-kind="{kind}"]')
        _click_reload(page, panel.get_by_role("button", name="Use this version").first)
    page.once("dialog", lambda dialog: dialog.accept())
    _click_reload(
        page,
        page.get_by_role("button", name="Confirm selected files — does not submit"),
    )


def _edited_docx_bytes(label: str) -> bytes:
    document = Document()
    document.add_heading(label, level=1)
    document.add_paragraph("Human-owned final wording edited outside JobSearch.")
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _edit_existing_docx_bytes(original: bytes, label: str) -> bytes:
    document = Document(BytesIO(original))
    document.add_paragraph(label)
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _assert_no_private_browser_content(page, live_server) -> None:
    html = page.content()
    visible = page.locator("body").inner_text()
    combined = html + visible
    assert live_server.secret not in combined
    assert "OPENAI_API_KEY" not in combined
    assert str(live_server.extensions_dir) not in combined
    assert "extension.json" not in combined


def test_browser_routes_do_not_expose_another_accounts_known_ids(page, live_server):
    conn = connect(live_server.db_path)
    create_account(conn, account_id="account_browser_b", display_name="Browser B")
    search_b = create_search_workspace(
        conn,
        account_id="account_browser_b",
        search_workspace_id="search_browser_b",
        name="Private Browser B Search",
    )
    application_b = create_workspace(
        conn,
        account_id="account_browser_b",
        workspace_id="ws_browser_b",
        company="Private Browser B Company",
        title="Private Browser B Role",
    )
    private_bytes = _edited_docx_bytes("Private Account B CV")
    blob = DocumentBlobStore(live_server.db_path.parent / "documents").publish(
        private_bytes
    )
    private_document = create_document_version(conn, {
        "id": "docv_browser_private_b", "account_id": "account_browser_b",
        "source_workspace_id": application_b["id"], "document_kind": "cv",
        "origin": "user_uploaded", "original_filename": "Private B.docx",
        "media_type": DOCX_MEDIA_TYPE, "byte_length": blob["byte_length"],
        "sha256": blob["sha256"], "storage_key": blob["storage_key"],
        "source_generation_artifact_id": None,
        "created_at": "2026-08-25T00:00:00+00:00",
    })
    conn.close()

    page.goto(live_server.base_url + "/search-workspaces")
    assert "Private Browser B Search" not in page.locator("body").inner_text()
    search_response = page.goto(
        live_server.base_url
        + f"/search-workspaces/{search_b['id']}/preferences"
    )
    assert search_response.status == 404
    application_response = page.goto(
        live_server.base_url + f"/workspaces/{application_b['id']}"
    )
    assert application_response.status == 404
    assert "Private Browser B Company" not in page.locator("body").inner_text()
    assert page.request.get(
        live_server.base_url
        + f"/api/workspaces/{application_b['id']}/application-documents/"
        + f"{private_document['id']}/download"
    ).status == 404


def test_user_profile_preferences_are_editable_in_browser(page, live_server):
    page.goto(f"{live_server.base_url}/user-profile", wait_until="networkidle")
    assert page.get_by_role("heading", name="Job search preferences").is_visible()
    assert page.get_by_text(
        "These preferences do not become candidate evidence and do not change Job Fit scoring."
    ).is_visible()

    page.locator('textarea[name="target_roles"]').fill("Project Manager\nProject Planner")
    page.locator('textarea[name="locations"]').fill("Aberdeen, UK\nRemote")
    page.locator('select[name="remote_preference"]').select_option("remote_or_hybrid")
    page.locator('input[name="recency_days"]').fill("30")
    page.locator('input[name="seniority_levels"][value="senior"]').check()
    page.locator('input[name="seniority_levels"][value="lead"]').check()
    page.locator('input[name="employment_types"][value="full_time"]').check()
    page.locator('textarea[name="industries"]').fill("Energy\nEngineering")
    page.locator('textarea[name="search_terms"]').fill("Primavera P6\nproject controls")
    page.locator('textarea[name="source_preferences"]').fill(
        "linkedin-search\nfreehire-search"
    )
    page.locator('input[name="compensation_currency"]').fill("GBP")
    page.locator('input[name="compensation_minimum"]').fill("60000")
    with page.expect_navigation(wait_until="networkidle"):
        page.get_by_role("button", name="Save search preferences").click()

    assert page.locator('textarea[name="target_roles"]').input_value() == (
        "Project Manager\nProject Planner"
    )
    assert page.locator('input[name="seniority_levels"][value="senior"]').is_checked()
    response = page.request.get(
        f"{live_server.base_url}/api/search-workspaces/search_default/user-profile"
    )
    assert response.status == 200
    payload = response.json()["user_profile"]["payload"]
    assert payload["target_roles"] == ["Project Manager", "Project Planner"]
    assert payload["compensation"] == {
        "currency": "GBP", "minimum": 60000, "period": "year",
    }
    _assert_no_private_browser_content(page, live_server)


def test_evidence_profile_manager_crud_sources_concurrency_and_staleness(page, live_server):
    _refresh_profile(page, live_server)
    workspace_url = _run_to_intelligence(page, live_server)

    page.goto(f"{live_server.base_url}/profile", wait_until="networkidle")
    assert page.get_by_role("heading", name="My Evidence Profile").is_visible()
    initial_revision = page.locator("#profile-manager").get_attribute("data-revision")
    conn = connect(live_server.db_path)
    initial_history = list_artifact_history(conn, PROFILE_WORKSPACE_ID, "profile_snapshot")
    conn.close()

    skill = page.locator(".profile-entry-card").filter(has_text="Python").first
    skill.get_by_role("button", name="Edit").click()
    skill.locator('input[name="value"]').fill("Python and SQL")
    with page.expect_navigation(wait_until="networkidle"):
        skill.get_by_role("button", name="Save").click()
    assert page.get_by_text("Python and SQL", exact=True).first.is_visible()

    location = page.locator(".profile-entry-card").filter(has_text="Location").first
    location.get_by_role("button", name="Edit").click()
    location.locator('input[name="value"]').fill("Birmingham, UK")
    with page.expect_navigation(wait_until="networkidle"):
        location.get_by_role("button", name="Save").click()
    assert page.get_by_role("heading", name="Conflicts detected").is_visible()
    assert page.get_by_text("Derived · read-only", exact=True).first.is_visible()
    assert page.get_by_role("button", name="Edit conflict").count() == 0

    page.get_by_text("Add profile information", exact=True).click()
    add_cert = page.locator('.profile-entry-add-form[data-entry-kind="certification"]')
    add_cert.locator('input[name="value"]').fill("PRINCE2 Practitioner")
    with page.expect_navigation(wait_until="networkidle"):
        add_cert.get_by_role("button", name="Add").click()
    certification = page.locator(".profile-entry-card").filter(has_text="PRINCE2 Practitioner").first
    assert certification.is_visible()

    certification.get_by_role("button", name="Edit").click()
    certification.locator('input[name="value"]').fill("Cancelled change")
    certification.get_by_role("button", name="Cancel").click()
    assert "PRINCE2 Practitioner" in certification.locator(".profile-entry-summary").inner_text()
    assert certification.locator(".profile-entry-edit-form").is_hidden()

    page.once("dialog", lambda dialog: dialog.accept())
    with page.expect_navigation(wait_until="networkidle"):
        certification.get_by_role("button", name="Delete").click()
    assert page.get_by_text("PRINCE2 Practitioner", exact=True).count() == 0

    assert page.locator('.profile-source[data-source-path="CLAUDE.md"]').count() == 0
    cv_source = page.locator('.profile-source[data-source-path="cv/main_example.tex"]')
    assert cv_source.get_by_text("Read-only", exact=True).is_visible()
    assert cv_source.get_by_role("button", name="Edit").count() == 0
    with page.expect_navigation(wait_until="networkidle"):
        cv_source.locator(".profile-source-toggle").uncheck()
    assert not page.locator('.profile-source[data-source-path="cv/main_example.tex"] .profile-source-toggle').is_checked()
    with page.expect_navigation(wait_until="networkidle"):
        page.locator('.profile-source[data-source-path="cv/main_example.tex"] .profile-source-toggle').check()

    stale_response = page.request.post(
        f"{live_server.base_url}/api/profile/entries",
        data={"expected_revision": initial_revision, "kind": "certification", "fields": {"value": "Stale write"}},
    )
    assert stale_response.status == 409
    assert "Reload it before saving" in stale_response.json()["detail"]

    conn = connect(live_server.db_path)
    history = list_artifact_history(conn, PROFILE_WORKSPACE_ID, "profile_snapshot")
    conn.close()
    assert len(history) >= len(initial_history) + 6
    assert history[-1]["payload"] == initial_history[-1]["payload"]

    page.goto(workspace_url, wait_until="networkidle")
    assert page.locator(".badge.stale").count() >= 1
    assert page.get_by_role("button", name="Rerun Job Fit").is_visible()
    _assert_no_private_browser_content(page, live_server)


def test_full_visible_journey_reaches_interview_with_explicit_submission(page, live_server):
    page.goto(live_server.base_url, wait_until="networkidle")
    assert page.get_by_role("heading", name="Your job pipeline").is_visible()
    assert page.get_by_text("No active applications").is_visible()
    _refresh_profile(page, live_server)
    workspace_url = _run_to_intelligence(page, live_server)

    assert page.locator('[data-item-id="cv-ready"]').count() == 2
    assert page.locator('[data-item-id="cv-needs-review"]').count() == 2
    unsupported = page.locator(".unsupported-record").first
    assert not unsupported.is_visible()
    assert page.get_by_text(
        "Technical details: resolved decisions and excluded claims", exact=True
    ).is_visible()
    assert unsupported.locator("button").count() == 0
    assert page.locator('[data-item-id="cv-unsupported-only"]').count() == 0
    assert page.get_by_role(
        "button", name="Use this"
    ).first.is_visible()
    assert page.get_by_role(
        "button", name="Leave this out"
    ).first.is_visible()
    assert page.locator("button.confirm-pack").is_disabled()
    assert page.locator('[data-status="drafted"]').count() == 0
    assert "applied" not in page.locator(".workspace-header").inner_text().casefold()
    _assert_no_private_browser_content(page, live_server)

    assert page.get_by_role("button", name="Use all generated text").is_visible()
    page.once("dialog", lambda dialog: dialog.accept())
    _click_reload(page, page.get_by_role("button", name="Use all generated text"))
    assert page.get_by_role("heading", name="Reviewed CV content").is_visible()
    assert page.locator("#reviewed-cv-content p").count() >= 2

    for _ in range(30):
        acknowledge = page.locator(
            'article.review-item:not(:has(.decision)) '
            'button.review-action[data-disposition="acknowledged_and_proceed"]'
        ).first
        if acknowledge.count() == 0:
            break
        _click_reload(page, acknowledge)
    else:
        raise AssertionError("review queue did not converge")
    assert page.get_by_text("0 outstanding", exact=True).is_visible()
    assert page.locator("button.confirm-pack").is_enabled()

    _confirm_pack(page)
    assert page.get_by_text("Workflow status:").locator("strong").inner_text() == "drafted"
    assert page.get_by_text("Generating or reviewing material never means it was submitted.").is_visible()
    assert page.get_by_role("heading", name="Is this application ready to send?").is_visible()
    assert page.get_by_text("Yes — ready to send", exact=True).is_visible()
    assert page.get_by_role("heading", name="Reviewed CV content").is_visible()
    assert page.get_by_role("heading", name="Reviewed cover letter content").is_visible()
    assert page.locator('[data-copy-section="cv"]').is_visible()
    assert page.locator('[data-copy-section="cover-letter"]').is_visible()

    cv_link = page.get_by_role("link", name="Download CV")
    cover_letter_link = page.get_by_role("link", name="Download Cover Letter")
    assert cv_link.is_visible()
    assert cover_letter_link.is_visible()
    cv_download = page.request.get(f"{live_server.base_url}{cv_link.get_attribute('href')}")
    cover_letter_download = page.request.get(
        f"{live_server.base_url}{cover_letter_link.get_attribute('href')}"
    )
    assert cv_download.status == 200
    assert cover_letter_download.status == 200
    assert cv_download.headers["content-type"] == (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    cv_document = Document(BytesIO(cv_download.body()))
    cv_texts = " ".join(p.text for p in cv_document.paragraphs)
    reviewed_cv_text = page.locator("#reviewed-cv-content").inner_text()
    assert reviewed_cv_text.strip().splitlines()[0].strip() in cv_texts

    cover_letter_document = Document(BytesIO(cover_letter_download.body()))
    cover_letter_texts = " ".join(p.text for p in cover_letter_document.paragraphs)
    assert cover_letter_texts.strip()
    assert "applied" not in page.locator(".workspace-header").inner_text().casefold()

    page.once("dialog", lambda dialog: dialog.accept())
    _click_reload(page, page.get_by_role("button", name="Mark applied — I submitted externally"))
    assert page.get_by_text("Workflow status:").locator("strong").inner_text() == "applied"
    assert page.get_by_role("link", name="Download CV").is_visible()
    assert page.get_by_role("link", name="Download Cover Letter").is_visible()

    _click_reload(page, page.get_by_role("button", name="Interview"))
    assert page.get_by_text("Workflow status:").locator("strong").inner_text() == "interview"
    assert page.get_by_role("link", name="Download CV").is_visible()
    assert page.get_by_role("link", name="Download Cover Letter").is_visible()

    page.goto(f"{live_server.base_url}/?filter=all", wait_until="networkidle")
    assert "active" in page.get_by_role("link", name="All").get_attribute("class")
    assert page.get_by_text("Browser Evidence Co").is_visible()
    assert page.get_by_text("Evidence Data Engineer").is_visible()
    assert page.get_by_text("interview", exact=True).is_visible()
    assert page.locator("th", has_text="Product stage").count() == 1
    assert page.locator("th", has_text="Application status").count() == 1
    assert page.locator("th", has_text="Next action").count() == 1
    page.get_by_role("link", name="Browser Evidence Co Evidence Data Engineer").click()
    assert page.url == workspace_url
    assert workspace_url.startswith(live_server.base_url + "/workspaces/")
    _assert_no_private_browser_content(page, live_server)


def test_application_pack_v1_embeds_candidate_facts_and_renders_history_exactly(
    page, live_server,
):
    candidate_path = (
        live_server.profile_root
        / ".claude/skills/job-application-assistant/01-candidate-profile.md"
    )
    candidate_path.write_text(
        candidate_path.read_text(encoding="utf-8")
        + """

## Education
- **MSc Computing** (2018-2020) - Example University — Key topics: Distributed systems

## Certifications
- **Cloud Professional**
""",
        encoding="utf-8",
    )
    _refresh_profile(page, live_server)
    workspace_url = _run_to_intelligence(page, live_server)
    _resolve_all_pending_reviews(page, "acknowledged_and_proceed")
    _confirm_pack(page)

    cv_link = page.get_by_role("link", name="Download CV")
    cover_link = page.get_by_role("link", name="Download Cover Letter")
    historical_cv_href = cv_link.get_attribute("href")
    historical_cover_href = cover_link.get_attribute("href")
    cv_before = page.request.get(f"{live_server.base_url}{historical_cv_href}")
    cover_before = page.request.get(f"{live_server.base_url}{historical_cover_href}")
    assert cv_before.status == cover_before.status == 200

    cv_texts = [
        paragraph.text
        for paragraph in Document(BytesIO(cv_before.body())).paragraphs
    ]
    cover_text = "\n".join(
        paragraph.text
        for paragraph in Document(BytesIO(cover_before.body())).paragraphs
    )
    assert cv_texts[0] == "Ada Lovelace"
    assert "Professional Experience" in cv_texts
    assert "Data Engineer | Evidence Works | 2020-01 - Present | London" in cv_texts
    assert "Built production data pipelines" in cv_texts
    assert "Tailored Highlights" in cv_texts
    assert cv_texts.index("Built production data pipelines") < cv_texts.index(
        "Tailored Highlights"
    )
    assert "MSc Computing" in "\n".join(cv_texts)
    assert "Cloud Professional" in cv_texts
    assert "Ada Lovelace" in cover_text
    assert "clm_9999999999999999" not in "\n".join(cv_texts) + cover_text

    candidate_path.write_text(
        candidate_path.read_text(encoding="utf-8")
        + "\n2. Ada Lovelace (2028). Must not enter the historical pack.\n",
        encoding="utf-8",
    )
    _refresh_profile(page, live_server)
    page.goto(workspace_url, wait_until="networkidle")
    assert page.get_by_role("link", name="Download CV").get_attribute(
        "href"
    ) == historical_cv_href

    cv_after = page.request.get(f"{live_server.base_url}{historical_cv_href}")
    cover_after = page.request.get(f"{live_server.base_url}{historical_cover_href}")
    assert cv_after.body() == cv_before.body()
    assert cover_after.body() == cover_before.body()
    assert cv_after.headers["x-content-hash"] == cv_before.headers["x-content-hash"]
    assert cover_after.headers["x-content-hash"] == cover_before.headers["x-content-hash"]
    historical_text = "\n".join(
        paragraph.text
        for paragraph in Document(BytesIO(cv_after.body())).paragraphs
    )
    assert "Must not enter the historical pack" not in historical_text
    _assert_no_private_browser_content(page, live_server)


def test_user_managed_documents_upload_select_confirm_replace_and_apply_exact_bytes(
    page, live_server,
):
    _refresh_profile(page, live_server)
    workspace_url = _run_to_intelligence(page, live_server)
    workspace_id = workspace_url.rsplit("/", 1)[-1]
    _resolve_all_pending_reviews(page, "acknowledged_and_proceed")

    _click_reload(page, page.get_by_role("button", name="Generate AI documents"))
    cv_panel = page.locator('[data-document-kind="cv"]')
    cover_panel = page.locator('[data-document-kind="cover_letter"]')
    assert cv_panel.get_by_text("No file selected").is_visible()
    assert cover_panel.get_by_text("No file selected").is_visible()

    cv_panel.locator('input[type="file"]').set_input_files({
        "name": "Not Word.docx", "mimeType": DOCX_MEDIA_TYPE,
        "buffer": b"not a ZIP package",
    })
    cv_panel.get_by_role("button", name="Upload").click()
    error_toast = page.locator("#toast.error").get_by_text(
        "file is not a DOCX ZIP package"
    )
    error_toast.wait_for(state="visible")
    assert error_toast.is_visible()
    assert cv_panel.locator(".document-version").filter(
        has_text="AI original"
    ).count() == 1

    edited = {
        "cv": ("Ada Final CV.docx", _edited_docx_bytes("Ada Final CV")),
        "cover_letter": (
            "Ada Final Cover Letter.docx",
            _edited_docx_bytes("Ada Final Cover Letter"),
        ),
    }
    for kind, panel in (("cv", cv_panel), ("cover_letter", cover_panel)):
        filename, content = edited[kind]
        panel.locator('input[type="file"]').set_input_files({
            "name": filename,
            "mimeType": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "buffer": content,
        })
        _click_reload(page, panel.get_by_role("button", name="Upload"))
        panel = page.locator(f'[data-document-kind="{kind}"]')
        assert panel.get_by_text("No file selected").is_visible()
        version = panel.locator(".document-version").filter(has_text=filename)
        assert version.get_by_text("not content-verified by JobSearch").is_visible()
        if kind == "cv":
            _click_reload(page, version.get_by_role("button", name="Save for reuse"))
            conn = connect(live_server.db_path)
            reusable_workspace = create_workspace(
                conn, company="Reuse Demo", title="Second Role"
            , account_id=DEFAULT_ACCOUNT_ID)
            conn.close()
            page.goto(
                f"{live_server.base_url}/workspaces/{reusable_workspace['id']}",
                wait_until="networkidle",
            )
            reusable = page.locator('[data-document-kind="cv"] .document-version').filter(
                has_text=filename
            )
            assert reusable.get_by_text(
                "not content-verified by JobSearch"
            ).is_visible()
            _click_reload(
                page, reusable.get_by_role("button", name="Use this version")
            )
            page.goto(workspace_url, wait_until="networkidle")
            panel = page.locator(f'[data-document-kind="{kind}"]')
            version = panel.locator(".document-version").filter(has_text=filename)
        _click_reload(page, version.get_by_role("button", name="Use this version"))

    page.once("dialog", lambda dialog: dialog.accept())
    _click_reload(
        page,
        page.get_by_role("button", name="Confirm selected files — does not submit"),
    )
    assert page.get_by_text("Workflow status:").locator("strong").inner_text() == "drafted"
    cv_download = page.request.get(
        f"{live_server.base_url}{page.get_by_role('link', name='Download CV').get_attribute('href')}"
    )
    cover_download = page.request.get(
        f"{live_server.base_url}{page.get_by_role('link', name='Download Cover Letter').get_attribute('href')}"
    )
    assert cv_download.body() == edited["cv"][1]
    assert cover_download.body() == edited["cover_letter"][1]

    conn = connect(live_server.db_path)
    first_pack = list_artifact_history(conn, workspace_id, "application_pack")[0]
    conn.close()
    first_cv_url = (
        f"{live_server.base_url}/api/workspaces/{workspace_id}/application-pack/render/cv"
        f"?pack_artifact_id={first_pack['id']}"
    )
    assert page.request.get(first_cv_url).body() == edited["cv"][1]

    cv_panel = page.locator('[data-document-kind="cv"]')
    ai_original = cv_panel.locator(".document-version").filter(has_text="AI original").first
    _click_reload(page, ai_original.get_by_role("button", name="Use this version"))
    assert page.get_by_text("Your selection has changed since confirmation").is_visible()
    assert page.request.get(first_cv_url).body() == edited["cv"][1]

    page.once("dialog", lambda dialog: dialog.accept())
    _click_reload(
        page,
        page.get_by_role("button", name="Confirm selected files — does not submit"),
    )
    assert page.request.get(first_cv_url).body() == edited["cv"][1]
    page.once("dialog", lambda dialog: dialog.accept())
    _click_reload(page, page.get_by_role("button", name="Mark applied — I submitted externally"))
    assert page.get_by_text("Workflow status:").locator("strong").inner_text() == "applied"
    assert page.locator(".document-upload-form").count() == 0
    assert page.get_by_role("button", name="Generate AI documents").is_disabled()
    _assert_no_private_browser_content(page, live_server)


def test_confirmed_pack_survives_upload_only_with_no_hidden_mutation_and_long_filename(
    page, live_server,
):
    _refresh_profile(page, live_server)
    workspace_url = _run_to_intelligence(page, live_server)
    workspace_id = workspace_url.rsplit("/", 1)[-1]
    _resolve_all_pending_reviews(page, "acknowledged_and_proceed")
    _click_reload(page, page.get_by_role("button", name="Generate AI documents"))

    for kind in ("cv", "cover_letter"):
        panel = page.locator(f'[data-document-kind="{kind}"]')
        original = panel.locator(".document-version").filter(has_text="AI original")
        _click_reload(page, original.get_by_role("button", name="Use this version"))

    page.once("dialog", lambda dialog: dialog.accept())
    _click_reload(
        page,
        page.get_by_role("button", name="Confirm selected files — does not submit"),
    )
    conn = connect(live_server.db_path)
    pack_a_id = get_current_artifact(conn, workspace_id, "application_pack")["id"]
    selection_before = get_selection(
        conn, workspace_id, "cv", account_id="account_local"
    )
    pack_count_before = conn.execute(
        "SELECT count(*) FROM artifacts WHERE workspace_id=? AND artifact_type='application_pack'",
        (workspace_id,),
    ).fetchone()[0]
    conn.close()

    cv_link = page.get_by_role("link", name="Download CV")
    pack_a_href = cv_link.get_attribute("href")
    assert f"pack_artifact_id={pack_a_id}" in pack_a_href
    assert f"pack_artifact_id={pack_a_id}" in page.get_by_role(
        "link", name="Download Cover Letter"
    ).get_attribute("href")
    original_download = page.request.get(f"{live_server.base_url}{pack_a_href}")
    assert original_download.status == 200
    selected_before_text = page.locator(
        '[data-document-kind="cv"] .selected-document'
    ).inner_text()

    long_filename = "Edited_" + ("deliberately_long_filename_" * 8) + "CV.docx"
    mutation_responses = []

    def capture_mutation(response):
        request = response.request
        if request.method in {"POST", "PUT", "PATCH", "DELETE"}:
            response_headers = response.headers
            mutation_responses.append({
                "captured_at_utc": datetime.now(timezone.utc).isoformat(),
                "method": request.method,
                "path": urlsplit(request.url).path,
                "status": response.status,
                "request_id": response_headers.get("x-request-id")
                or response_headers.get("x-correlation-id"),
            })

    page.on("response", capture_mutation)
    cv_panel = page.locator('[data-document-kind="cv"]')
    cv_panel.locator('input[type="file"]').set_input_files({
        "name": long_filename,
        "mimeType": DOCX_MEDIA_TYPE,
        "buffer": _edit_existing_docx_bytes(
            original_download.body(),
            "Uploaded replacement must remain unselected",
        ),
    })
    _click_reload(page, cv_panel.get_by_role("button", name="Upload"))
    page.remove_listener("response", capture_mutation)

    upload_path = f"/api/workspaces/{workspace_id}/application-documents/upload/cv"
    selection_path = f"/api/workspaces/{workspace_id}/application-documents/selection/cv"
    pack_path = f"/api/workspaces/{workspace_id}/application-pack"
    assert sum(
        item["method"] == "POST" and item["path"] == upload_path
        for item in mutation_responses
    ) == 1
    assert not any(
        item["method"] == "PUT" and item["path"] == selection_path
        for item in mutation_responses
    )
    assert not any(
        item["method"] == "POST" and item["path"] == pack_path
        for item in mutation_responses
    )

    conn = connect(live_server.db_path)
    selection_after = get_selection(
        conn, workspace_id, "cv", account_id="account_local"
    )
    assert selection_after == selection_before
    assert get_current_artifact(conn, workspace_id, "application_pack")["id"] == pack_a_id
    assert conn.execute(
        "SELECT count(*) FROM artifacts WHERE workspace_id=? AND artifact_type='application_pack'",
        (workspace_id,),
    ).fetchone()[0] == pack_count_before
    conn.close()

    assert page.get_by_role("link", name="Download CV").get_attribute("href") == pack_a_href
    after_upload_download = page.request.get(f"{live_server.base_url}{pack_a_href}")
    assert after_upload_download.body() == original_download.body()
    assert after_upload_download.headers["x-content-hash"] == original_download.headers[
        "x-content-hash"
    ]
    cv_panel = page.locator('[data-document-kind="cv"]')
    selected_text = cv_panel.locator(".selected-document").inner_text()
    assert selected_text == selected_before_text
    assert f"revision {selection_before['revision']}" in selected_text
    replacement = cv_panel.locator(".document-version").filter(has_text=long_filename)
    assert replacement.get_by_text("not content-verified by JobSearch").is_visible()

    print("FRESH_ACCEPTANCE_TRACE " + json.dumps({
        "workspace_id": workspace_id,
        "pack_artifact_id": pack_a_id,
        "selection_revision": selection_after["revision"],
        "selected_document_version_id": selection_after["document_version_id"],
        "confirmed_cv_sha256": original_download.headers["x-content-hash"],
        "mutation_responses_after_upload_started": mutation_responses,
    }, sort_keys=True))

    page.set_viewport_size({"width": 480, "height": 900})
    layout = replacement.evaluate("""element => {
      const panel = element.closest('.document-kind-panel');
      const panelBox = panel.getBoundingClientRect();
      const controls = [...element.querySelectorAll('a, button')].map(control => {
        const box = control.getBoundingClientRect();
        return {left: box.left, right: box.right};
      });
      return {
        panelLeft: panelBox.left,
        panelRight: panelBox.right,
        scrollWidth: element.scrollWidth,
        clientWidth: element.clientWidth,
        controls,
      };
    }""")
    assert layout["scrollWidth"] <= layout["clientWidth"] + 1
    assert all(
        layout["panelLeft"] - 1 <= control["left"]
        and control["right"] <= layout["panelRight"] + 1
        for control in layout["controls"]
    )


def test_stale_and_review_negative_paths_are_enforced_in_rendered_ui(page, live_server):
    _refresh_profile(page, live_server)
    workspace_url = _run_to_intelligence(page, live_server)
    assert page.locator("button.confirm-pack").is_disabled()
    assert page.locator('[data-item-id="cv-needs-review"]').count() == 2
    assert page.locator('[data-status="drafted"]').count() == 0
    assert "applied" not in page.locator(".workspace-header").inner_text().casefold()

    for _ in range(30):
        omit = page.locator(
            'article.review-item:not(:has(.decision)) '
            'button.review-action[data-disposition="omit_from_positioning"]'
        ).first
        if omit.count() == 0:
            break
        _click_reload(page, omit)
    else:
        raise AssertionError("omission review queue did not converge")
    assert page.get_by_text("0 outstanding", exact=True).is_visible()
    assert page.get_by_text("INCOMPLETE", exact=True).is_visible()
    assert page.locator("button.confirm-pack").is_disabled()

    workspace_id = workspace_url.rsplit("/", 1)[-1]
    pack_response = page.request.post(
        f"{live_server.base_url}/api/workspaces/{workspace_id}/application-pack",
        data={"confirmed": True, "effective_date": "2026-08-21"},
    )
    assert pack_response.status == 400
    assert "no reviewed usable application material" in pack_response.json()["detail"]
    applied_response = page.request.patch(
        f"{live_server.base_url}/api/workspaces/{workspace_id}/status",
        data={"new_status": "applied", "effective_date": "2026-08-21"},
    )
    assert applied_response.status == 400
    assert page.locator('[data-status="drafted"]').count() == 0

    candidate_path = (
        live_server.profile_root
        / ".claude/skills/job-application-assistant/01-candidate-profile.md"
    )
    candidate_path.write_text(
        candidate_path.read_text(encoding="utf-8")
        + "\n2. Ada Lovelace (2027). A new browser-staleness publication.\n",
        encoding="utf-8",
    )
    _refresh_profile(page, live_server)
    page.goto(workspace_url, wait_until="networkidle")
    assert page.locator(".badge.stale").count() >= 1
    assert page.locator("button.confirm-pack").is_disabled()
    assert page.get_by_role("button", name="Rerun Job Fit").is_visible()
    assert page.get_by_role("button", name="Rerun Application Intelligence").count() == 0
    assert page.locator('[data-status="drafted"]').count() == 0
    _assert_no_private_browser_content(page, live_server)
    page.locator('input[name="extension_ids"][value="data-transfer"]').check()
    _click_reload(page, page.get_by_role("button", name="Rerun Job Fit"))
    assert page.get_by_role("button", name="Rerun Application Intelligence").is_visible()
    assert page.locator("button.confirm-pack").is_disabled()
    _click_reload(page, page.get_by_role("button", name="Rerun Application Intelligence"))
    assert page.locator('[data-item-id="cv-ready"]').count() == 2
    assert page.locator('[data-item-id="cv-needs-review"]').count() == 2
    assert page.locator("button.confirm-pack").is_disabled()

    for _ in range(30):
        acknowledge = page.locator(
            'article.review-item:not(:has(.decision)) '
            'button.review-action[data-disposition="acknowledged_and_proceed"]'
        ).first
        if acknowledge.count() == 0:
            break
        _click_reload(page, acknowledge)
    else:
        raise AssertionError("recovered review queue did not converge")
    assert page.get_by_text("0 outstanding", exact=True).is_visible()
    assert page.locator("button.confirm-pack").is_enabled()
    _assert_no_private_browser_content(page, live_server)


def test_causal_staleness_message_appears_and_differs_by_stage(page, live_server):
    _refresh_profile(page, live_server)
    workspace_url = _run_to_intelligence(page, live_server)

    candidate_path = (
        live_server.profile_root
        / ".claude/skills/job-application-assistant/01-candidate-profile.md"
    )
    candidate_path.write_text(
        candidate_path.read_text(encoding="utf-8")
        + "\n2. Ada Lovelace (2027). A new browser-causal-staleness publication.\n",
        encoding="utf-8",
    )
    _refresh_profile(page, live_server)
    page.goto(workspace_url, wait_until="networkidle")

    fit_panel = page.locator("#job-fit")
    assert fit_panel.get_by_text("Evidence Profile").is_visible()
    assert fit_panel.get_by_role("button", name="Rerun Job Fit").is_visible()

    page.locator('input[name="extension_ids"][value="data-transfer"]').check()
    _click_reload(page, page.get_by_role("button", name="Rerun Job Fit"))

    intelligence_panel = page.locator("#application-intelligence")
    assert intelligence_panel.get_by_text("Job Fit").is_visible()
    assert intelligence_panel.get_by_role(
        "button", name="Rerun Application Intelligence"
    ).is_visible()
    _assert_no_private_browser_content(page, live_server)


def test_gate_four_reason_survives_an_existing_confirmed_pack(page, live_server):
    _refresh_profile(page, live_server)
    workspace_url = _run_to_intelligence(page, live_server)

    _resolve_all_pending_reviews(page, "acknowledged_and_proceed")
    _confirm_pack(page)
    assert page.get_by_text("Yes — ready to send", exact=True).is_visible()

    candidate_path = (
        live_server.profile_root
        / ".claude/skills/job-application-assistant/01-candidate-profile.md"
    )
    candidate_path.write_text(
        candidate_path.read_text(encoding="utf-8")
        + "\n2. Ada Lovelace (2027). A new browser-gate-four publication.\n",
        encoding="utf-8",
    )
    _refresh_profile(page, live_server)
    page.goto(workspace_url, wait_until="networkidle")
    page.locator('input[name="extension_ids"][value="data-transfer"]').check()
    _click_reload(page, page.get_by_role("button", name="Rerun Job Fit"))
    _click_reload(
        page, page.get_by_role("button", name="Rerun Application Intelligence")
    )

    _resolve_all_pending_reviews(page, "omit_from_positioning")

    assert page.get_by_text("INCOMPLETE", exact=True).is_visible()
    assert page.locator("button.confirm-pack").is_disabled()
    assert page.get_by_text(
        "Reviewed material must be completion-ready before AI documents can be generated."
    ).is_visible()
    assert page.get_by_text(
        "This document was generated from earlier reviewed material."
    ).first.is_visible()
    _assert_no_private_browser_content(page, live_server)


def test_friendly_completion_counts_visible_when_material_incomplete(page, live_server):
    _refresh_profile(page, live_server)
    _run_to_intelligence(page, live_server)

    _resolve_all_pending_reviews(page, "omit_from_positioning")

    assert page.get_by_text("INCOMPLETE", exact=True).is_visible()
    assert page.get_by_text("0 of 2 required CV bullets").first.is_visible()
    assert page.get_by_text("No decisions are pending, but the application material is not yet usable.").is_visible()
    assert page.get_by_text("No decisions need your attention.", exact=True).count() == 0
    assert page.get_by_role("link", name="Review your profile").is_visible()
    assert page.get_by_role("link", name="Review job evidence").is_visible()
    assert page.get_by_role("button", name="Rerun Application Intelligence").is_visible()
    assert page.get_by_text("insufficient_cv_units", exact=True).count() == 0
    assert page.get_by_text("—", exact=False).count() > 0


def test_how_it_works_page_reachable_from_nav_with_pipeline_and_glossary(
    page, live_server
):
    page.goto(live_server.base_url, wait_until="networkidle")
    page.get_by_role("link", name="How it works", exact=True).click()
    page.wait_for_url("**/how-it-works")
    assert page.get_by_role("heading", name="How this app works").is_visible()
    for stage_name in (
        "Evidence Profile", "Find/Add Job", "Understanding", "Job Fit",
        "Application Intelligence", "Application Pack",
    ):
        assert page.get_by_text(stage_name).first.is_visible()
    for term in (
        "Current", "Stale / needs updating", "Needs review", "Blocked / incomplete",
        "Ready", "Historical pack", "Drafted", "Applied",
    ):
        assert page.get_by_text(term, exact=True).first.is_visible()
    _assert_no_private_browser_content(page, live_server)


def test_getting_started_card_visible_on_dashboard_and_links_to_how_it_works(
    page, live_server
):
    page.goto(live_server.base_url, wait_until="networkidle")
    card = page.locator(".getting-started-card")
    assert card.is_visible()
    for label in (
        "Evidence Profile", "Find/Add Job", "Job Fit", "Intelligence",
        "Review", "Pack", "Download", "Apply",
    ):
        assert card.get_by_text(label, exact=True).is_visible()
    with page.expect_navigation(wait_until="networkidle"):
        card.get_by_role("link", name="See the full walkthrough").click()
    assert page.url.endswith("/how-it-works")


def test_reviewed_output_empty_state_names_next_action(page, live_server):
    _refresh_profile(page, live_server)
    _run_to_intelligence(page, live_server)
    empty_state = page.locator("#reviewed-cv-content .muted")
    assert empty_state.is_visible()
    text = empty_state.inner_text()
    assert "Resolve" in text or "Application Intelligence" in text


def test_historical_pack_and_incomplete_current_material_never_read_as_contradictory(
    page, live_server,
):
    _refresh_profile(page, live_server)
    workspace_url = _run_to_intelligence(page, live_server)

    _resolve_all_pending_reviews(page, "acknowledged_and_proceed")
    _confirm_pack(page)
    cv_link = page.get_by_role("link", name="Download CV")
    cover_letter_link = page.get_by_role("link", name="Download Cover Letter")
    assert cv_link.is_visible()
    assert cover_letter_link.is_visible()
    historical_cv_href = cv_link.get_attribute("href")
    historical_cover_href = cover_letter_link.get_attribute("href")

    candidate_path = (
        live_server.profile_root
        / ".claude/skills/job-application-assistant/01-candidate-profile.md"
    )
    candidate_path.write_text(
        candidate_path.read_text(encoding="utf-8")
        + "\n2. Ada Lovelace (2027). A new browser-combined-regression publication.\n",
        encoding="utf-8",
    )
    _refresh_profile(page, live_server)
    page.goto(workspace_url, wait_until="networkidle")
    assert page.locator(".badge.stale").count() >= 1
    page.locator('input[name="extension_ids"][value="data-transfer"]').check()
    _click_reload(page, page.get_by_role("button", name="Rerun Job Fit"))
    _click_reload(
        page, page.get_by_role("button", name="Rerun Application Intelligence")
    )

    _resolve_all_pending_reviews(page, "omit_from_positioning")

    # 1. Historical downloads remain available and point at the same rendered artifact.
    cv_link = page.get_by_role("link", name="Download CV")
    cover_letter_link = page.get_by_role("link", name="Download Cover Letter")
    assert cv_link.is_visible()
    assert cover_letter_link.is_visible()
    assert cv_link.get_attribute("href") == historical_cv_href
    assert cover_letter_link.get_attribute("href") == historical_cover_href
    cv_download = page.request.get(f"{live_server.base_url}{historical_cv_href}")
    assert cv_download.status == 200

    # 2. The historical pack is explicitly labeled as previous/confirmed, not
    #    presented as if it were the freshly reviewed material.
    assert page.get_by_text(
        "Confirmed files remain immutable and downloadable:"
    ).is_visible()

    # 3. A separate, distinct statement says the replacement pack is not ready,
    #    with the actual current completion issue visible.
    assert page.get_by_text(
        "Reviewed material must be completion-ready before AI documents can be generated."
    ).is_visible()
    assert page.get_by_text("INCOMPLETE", exact=True).is_visible()
    assert page.get_by_text("required CV bullets").first.is_visible()

    # 4. The confirm-pack button stays disabled — no automatic replacement.
    assert page.locator("button.confirm-pack").is_disabled()

    # 5. The reviewed-content panel heading is qualified as historical, never
    #    presented as the current unreviewed material.
    assert page.get_by_text(
        "Reviewed CV content (from your confirmed pack)"
    ).is_visible()

    _assert_no_private_browser_content(page, live_server)


def test_discovery_search_evaluate_and_promote_browser_lifecycle(page, live_server):
    _refresh_profile(page, live_server)
    page.goto(f"{live_server.base_url}/user-profile", wait_until="networkidle")
    page.locator('textarea[name="target_roles"]').fill("Data Engineer")
    page.locator('textarea[name="locations"]').fill("London, UK")
    page.locator('textarea[name="search_terms"]').fill("Python data pipelines")
    page.locator('textarea[name="source_preferences"]').fill("freehire-search")
    _click_reload(page, page.get_by_role("button", name="Save search preferences"))

    page.goto(f"{live_server.base_url}/discover", wait_until="networkidle")
    assert page.get_by_role("heading", name="Discover and rank jobs").is_visible()
    assert page.get_by_text("Adjusting this search does not change Job Fit scoring.").is_visible()
    _click_reload(page, page.get_by_role("button", name="Search jobs"))

    card = page.locator('[data-candidate-id]').filter(has_text="Discovery Evidence Co")
    assert card.get_by_text("Evidence Data Engineer").is_visible()
    assert card.get_by_text("No invented score").is_visible()
    card.locator(".candidate-select").check()
    _click_reload(page, page.get_by_role("button", name="Evaluate selected").first)

    card = page.locator('[data-candidate-id]').filter(has_text="Discovery Evidence Co")
    assert card.get_by_text("No invented score").is_visible()
    _click_reload(page, card.get_by_role("button", name="Save"))
    card = page.locator('[data-candidate-id]').filter(has_text="Discovery Evidence Co")
    with page.expect_navigation(wait_until="networkidle"):
        card.get_by_role("button", name="Create application").click()
    assert "/workspaces/" in page.url
    assert page.get_by_text("Discovery Evidence Co", exact=True).is_visible()
    assert page.get_by_role("heading", name="Evidence Data Engineer").is_visible()

    page.goto(f"{live_server.base_url}/discover", wait_until="networkidle")
    assert page.locator('[data-candidate-id]').filter(has_text="Discovery Evidence Co").count() == 1
    applications = page.request.get(f"{live_server.base_url}/api/workspaces").json()["workspaces"]
    assert len([item for item in applications if item["company"] == "Discovery Evidence Co"]) == 1
    _assert_no_private_browser_content(page, live_server)


def test_search_workspace_switching_keeps_preferences_isolated(page, live_server):
    page.goto(f"{live_server.base_url}/user-profile", wait_until="networkidle")
    page.locator('textarea[name="target_roles"]').fill("Project Planner")
    _click_reload(page, page.get_by_role("button", name="Save search preferences"))

    page.goto(f"{live_server.base_url}/search-workspaces", wait_until="networkidle")
    page.locator('#create-search-workspace-form input[name="name"]').fill("Project Manager")
    with page.expect_navigation(wait_until="networkidle"):
        page.get_by_role("button", name="Create workspace").click()
    assert "/preferences" in page.url
    assert page.get_by_text("Project Manager · Search preferences").is_visible()
    page.locator('textarea[name="target_roles"]').fill("Project Manager")
    _click_reload(page, page.get_by_role("button", name="Save search preferences"))

    page.get_by_label("Search workspace", exact=True).select_option(label="Default search")
    page.wait_for_url("**/search-workspaces/search_default/discover")
    page.get_by_role("link", name="Search preferences", exact=True).click()
    page.wait_for_url("**/search-workspaces/search_default/preferences")
    assert page.locator('textarea[name="target_roles"]').input_value() == "Project Planner"
    _assert_no_private_browser_content(page, live_server)
