from __future__ import annotations

from webapp.persistence.accounts import create_account
from webapp.persistence.db import connect, init_db
from webapp.services.handoff import (
    PairingSecretInvalid,
)


def _conn(tmp_path):
    db_path = tmp_path / "jobsearch.sqlite3"
    init_db(db_path)
    return connect(db_path)


from datetime import datetime, timedelta, timezone

from webapp.services.handoff import PairingSecretExpired


from webapp.persistence.artifacts import save_artifact
from webapp.persistence.workspaces import create_workspace, ensure_profile_workspace
from webapp.services.ownership import AccountScope, OwnedResourceNotFound, account_profile_root
from webapp.services.handoff import (
    HandoffPackNotFound,
    discover_resumable_handoff_sessions,
    start_handoff_session,
)


def _scope(tmp_path):
    return AccountScope(
        account_id="account_local",
        profile_root=account_profile_root(str(tmp_path), "account_local"),
    )


def _workspace_with_pack(conn):
    ensure_profile_workspace(conn, account_id="account_local")
    workspace = create_workspace(
        conn, company="Acme", title="Engineer", account_id="account_local",
    )
    artifact = save_artifact(
        conn, workspace_id=workspace["id"], artifact_type="application_pack",
        payload={"schema_version": "application-pack.v1"},
    )
    return workspace, artifact


def test_start_handoff_session_succeeds_for_owned_workspace_and_pack(tmp_path):
    conn = _conn(tmp_path)
    scope = _scope(tmp_path)
    workspace, artifact = _workspace_with_pack(conn)

    session = start_handoff_session(
        conn, scope, workspace_id=workspace["id"], pack_artifact_id=artifact["id"],
        target_url="https://boards.greenhouse.io/acme/jobs/1",
        target_domain="boards.greenhouse.io", ats_adapter_id="greenhouse",
        ats_adapter_version="greenhouse@1",
    )
    assert session["workspace_id"] == workspace["id"]
    assert session["pack_artifact_id"] == artifact["id"]
    assert session["account_id"] == "account_local"
    conn.close()


def test_start_handoff_session_rejects_workspace_not_owned_by_scope(tmp_path):
    from webapp.persistence.accounts import create_account

    conn = _conn(tmp_path)
    create_account(conn, account_id="account_other", display_name="Other")
    workspace, artifact = _workspace_with_pack(conn)

    other_scope = AccountScope(
        account_id="account_other",
        profile_root=account_profile_root(str(tmp_path), "account_other"),
    )
    try:
        start_handoff_session(
            conn, other_scope, workspace_id=workspace["id"],
            pack_artifact_id=artifact["id"], target_url="https://x.test/apply",
            target_domain="x.test", ats_adapter_id="generic",
            ats_adapter_version="generic@1",
        )
        assert False, "expected OwnedResourceNotFound"
    except OwnedResourceNotFound:
        pass
    conn.close()


def test_start_handoff_session_rejects_pack_from_different_workspace(tmp_path):
    conn = _conn(tmp_path)
    scope = _scope(tmp_path)
    workspace, _artifact = _workspace_with_pack(conn)
    other_workspace = create_workspace(
        conn, company="Other Co", title="Role", account_id="account_local",
    )
    foreign_artifact = save_artifact(
        conn, workspace_id=other_workspace["id"], artifact_type="application_pack",
        payload={"schema_version": "application-pack.v1"},
    )

    try:
        start_handoff_session(
            conn, scope, workspace_id=workspace["id"],
            pack_artifact_id=foreign_artifact["id"],
            target_url="https://x.test/apply", target_domain="x.test",
            ats_adapter_id="generic", ats_adapter_version="generic@1",
        )
        assert False, "expected HandoffPackNotFound"
    except HandoffPackNotFound:
        pass
    conn.close()


def test_discover_resumable_sessions_scoped_to_owned_workspace(tmp_path):
    conn = _conn(tmp_path)
    scope = _scope(tmp_path)
    workspace, artifact = _workspace_with_pack(conn)
    start_handoff_session(
        conn, scope, workspace_id=workspace["id"], pack_artifact_id=artifact["id"],
        target_url="https://x.test/apply", target_domain="x.test",
        ats_adapter_id="generic", ats_adapter_version="generic@1",
    )

    found = discover_resumable_handoff_sessions(
        conn, scope, workspace_id=workspace["id"], target_domain="x.test",
    )
    assert len(found) == 1
    conn.close()


from webapp.services.handoff import (
    HandoffEventRejected,
    HandoffSessionNotActive,
    HandoffSessionNotFound,
    mint_session_token,
    record_handoff_event,
    replay_handoff_session,
    resolve_session_scope,
    resume_handoff_session,
)


def _session_scope(conn, session):
    """Mints a real session token for an already-started session and
    resolves it into the SessionScope record_handoff_event/
    replay_handoff_session/confirm_handoff_submission now require —
    matching how the API layer actually obtains one via
    get_session_scope, rather than constructing a SessionScope by hand."""
    token = mint_session_token(conn, handoff_session_id=session["id"])
    return resolve_session_scope(conn, raw_token=token)


def test_record_handoff_event_succeeds_for_owned_in_progress_session(tmp_path):
    conn = _conn(tmp_path)
    scope = _scope(tmp_path)
    workspace, artifact = _workspace_with_pack(conn)
    session = start_handoff_session(
        conn, scope, workspace_id=workspace["id"], pack_artifact_id=artifact["id"],
        target_url="https://x.test/apply", target_domain="x.test",
        ats_adapter_id="generic", ats_adapter_version="generic@1",
    )
    session_scope = _session_scope(conn, session)

    event = record_handoff_event(
        conn, session_scope, handoff_session_id=session["id"], event_id="evt_1",
        event_type="value_inserted",
        event_payload={"value": "shola@example.com"},
        normalized_field_type="email", page_field_key="generic:email",
    )
    assert event["event_type"] == "value_inserted"
    conn.close()


def test_record_handoff_event_rejects_value_on_presence_only_event_type(tmp_path):
    conn = _conn(tmp_path)
    scope = _scope(tmp_path)
    workspace, artifact = _workspace_with_pack(conn)
    session = start_handoff_session(
        conn, scope, workspace_id=workspace["id"], pack_artifact_id=artifact["id"],
        target_url="https://x.test/apply", target_domain="x.test",
        ats_adapter_id="generic", ats_adapter_version="generic@1",
    )
    session_scope = _session_scope(conn, session)

    try:
        record_handoff_event(
            conn, session_scope, handoff_session_id=session["id"], event_id="evt_sensitive",
            event_type="user_value_present_observed",
            event_payload={"value": "should not be here"},
            normalized_field_type="salary_expectation",
            page_field_key="generic:salary",
        )
        assert False, "expected HandoffEventRejected"
    except HandoffEventRejected:
        pass
    conn.close()


def test_record_handoff_event_rejects_session_not_owned_by_scope(tmp_path):
    conn = _conn(tmp_path)
    scope = _scope(tmp_path)
    workspace, artifact = _workspace_with_pack(conn)
    session = start_handoff_session(
        conn, scope, workspace_id=workspace["id"], pack_artifact_id=artifact["id"],
        target_url="https://x.test/apply", target_domain="x.test",
        ats_adapter_id="generic", ats_adapter_version="generic@1",
    )

    # A session token is bound to exactly one handoff_session_id — a
    # token minted for a DIFFERENT session must never work against this
    # one, even for the same account. Model "not owned by scope" as
    # exactly that, since SessionScope no longer carries a bare account
    # identity a caller could present against an arbitrary session id.
    other_session = start_handoff_session(
        conn, scope, workspace_id=workspace["id"], pack_artifact_id=artifact["id"],
        target_url="https://other.test/apply", target_domain="other.test",
        ats_adapter_id="generic", ats_adapter_version="generic@1",
    )
    other_session_scope = _session_scope(conn, other_session)

    try:
        record_handoff_event(
            conn, other_session_scope, handoff_session_id=session["id"], event_id="evt_x",
            event_type="field_detected", event_payload={},
        )
        assert False, "expected HandoffSessionNotFound"
    except HandoffSessionNotFound:
        pass
    conn.close()


def test_record_handoff_event_rejects_terminal_session(tmp_path):
    from webapp.persistence.handoff import set_handoff_session_status

    conn = _conn(tmp_path)
    scope = _scope(tmp_path)
    workspace, artifact = _workspace_with_pack(conn)
    session = start_handoff_session(
        conn, scope, workspace_id=workspace["id"], pack_artifact_id=artifact["id"],
        target_url="https://x.test/apply", target_domain="x.test",
        ats_adapter_id="generic", ats_adapter_version="generic@1",
    )
    session_scope = _session_scope(conn, session)
    set_handoff_session_status(conn, session["id"], status="expired")

    try:
        record_handoff_event(
            conn, session_scope, handoff_session_id=session["id"], event_id="evt_late",
            event_type="field_detected", event_payload={},
        )
        assert False, "expected HandoffSessionNotActive"
    except HandoffSessionNotActive:
        pass
    conn.close()


def test_replay_handoff_session_returns_session_and_ordered_events(tmp_path):
    conn = _conn(tmp_path)
    scope = _scope(tmp_path)
    workspace, artifact = _workspace_with_pack(conn)
    session = start_handoff_session(
        conn, scope, workspace_id=workspace["id"], pack_artifact_id=artifact["id"],
        target_url="https://x.test/apply", target_domain="x.test",
        ats_adapter_id="generic", ats_adapter_version="generic@1",
    )
    session_scope = _session_scope(conn, session)
    record_handoff_event(
        conn, session_scope, handoff_session_id=session["id"], event_id="evt_1",
        event_type="field_detected", event_payload={},
    )
    record_handoff_event(
        conn, session_scope, handoff_session_id=session["id"], event_id="evt_2",
        event_type="value_inserted", event_payload={"value": "x"},
    )

    replay = replay_handoff_session(conn, session_scope, session["id"])
    assert replay["session"]["id"] == session["id"]
    assert [e["event_id"] for e in replay["events"]] == ["evt_1", "evt_2"]
    conn.close()


def test_replay_handoff_session_rejects_session_not_owned_by_scope(tmp_path):
    conn = _conn(tmp_path)
    scope = _scope(tmp_path)
    workspace, artifact = _workspace_with_pack(conn)
    session = start_handoff_session(
        conn, scope, workspace_id=workspace["id"], pack_artifact_id=artifact["id"],
        target_url="https://x.test/apply", target_domain="x.test",
        ats_adapter_id="generic", ats_adapter_version="generic@1",
    )

    other_session = start_handoff_session(
        conn, scope, workspace_id=workspace["id"], pack_artifact_id=artifact["id"],
        target_url="https://other.test/apply", target_domain="other.test",
        ats_adapter_id="generic", ats_adapter_version="generic@1",
    )
    other_session_scope = _session_scope(conn, other_session)

    try:
        replay_handoff_session(conn, other_session_scope, session["id"])
        assert False, "expected HandoffSessionNotFound"
    except HandoffSessionNotFound:
        pass
    conn.close()


from webapp.persistence.workflow import record_status_change
from webapp.services.handoff import confirm_handoff_submission


def test_confirm_handoff_submission_without_workflow_update(tmp_path):
    conn = _conn(tmp_path)
    scope = _scope(tmp_path)
    workspace, artifact = _workspace_with_pack(conn)
    session = start_handoff_session(
        conn, scope, workspace_id=workspace["id"], pack_artifact_id=artifact["id"],
        target_url="https://x.test/apply", target_domain="x.test",
        ats_adapter_id="generic", ats_adapter_version="generic@1",
    )
    session_scope = _session_scope(conn, session)

    result = confirm_handoff_submission(
        conn, session_scope, handoff_session_id=session["id"], mark_workflow_applied=False,
    )
    assert result["session"]["status"] == "user_confirmed_submitted"
    assert result["confirmation"]["handoff_session_id"] == session["id"]
    assert result["workflow_event"] is None
    conn.close()


def test_confirm_handoff_submission_rejects_session_not_owned(tmp_path):
    conn = _conn(tmp_path)
    scope = _scope(tmp_path)
    workspace, artifact = _workspace_with_pack(conn)
    session = start_handoff_session(
        conn, scope, workspace_id=workspace["id"], pack_artifact_id=artifact["id"],
        target_url="https://x.test/apply", target_domain="x.test",
        ats_adapter_id="generic", ats_adapter_version="generic@1",
    )

    other_session = start_handoff_session(
        conn, scope, workspace_id=workspace["id"], pack_artifact_id=artifact["id"],
        target_url="https://other.test/apply", target_domain="other.test",
        ats_adapter_id="generic", ats_adapter_version="generic@1",
    )
    other_session_scope = _session_scope(conn, other_session)

    try:
        confirm_handoff_submission(
            conn, other_session_scope, handoff_session_id=session["id"],
        )
        assert False, "expected HandoffSessionNotFound"
    except HandoffSessionNotFound:
        pass
    conn.close()


def test_confirming_twice_is_rejected_not_double_recorded(tmp_path):
    conn = _conn(tmp_path)
    scope = _scope(tmp_path)
    workspace, artifact = _workspace_with_pack(conn)
    session = start_handoff_session(
        conn, scope, workspace_id=workspace["id"], pack_artifact_id=artifact["id"],
        target_url="https://x.test/apply", target_domain="x.test",
        ats_adapter_id="generic", ats_adapter_version="generic@1",
    )
    session_scope = _session_scope(conn, session)
    confirm_handoff_submission(conn, session_scope, handoff_session_id=session["id"])

    try:
        confirm_handoff_submission(conn, session_scope, handoff_session_id=session["id"])
        assert False, "expected HandoffSessionNotActive"
    except HandoffSessionNotActive:
        pass
    conn.close()


from datetime import datetime, timedelta, timezone

from webapp.persistence.handoff import hash_pairing_secret
from webapp.services.handoff import (
    HandoffSessionExpired,
    HandoffSessionTokenInvalid,
    SessionScope,
    mint_session_token,
    resolve_session_scope,
    rotate_session_token,
)


def test_mint_session_token_persists_hash_not_raw_value(tmp_path):
    conn = _conn(tmp_path)
    scope = _scope(tmp_path)
    workspace, artifact = _workspace_with_pack(conn)
    session = start_handoff_session(
        conn, scope, workspace_id=workspace["id"], pack_artifact_id=artifact["id"],
        target_url="https://x.test/apply", target_domain="x.test",
        ats_adapter_id="generic", ats_adapter_version="generic@1",
    )

    token = mint_session_token(conn, handoff_session_id=session["id"])

    row = conn.execute(
        "SELECT token_hash FROM handoff_session_tokens WHERE handoff_session_id = ?",
        (session["id"],),
    ).fetchone()
    assert row["token_hash"] != token
    assert row["token_hash"] == hash_pairing_secret(token)
    conn.close()


def test_resolve_session_scope_succeeds_for_valid_token(tmp_path):
    conn = _conn(tmp_path)
    scope = _scope(tmp_path)
    workspace, artifact = _workspace_with_pack(conn)
    session = start_handoff_session(
        conn, scope, workspace_id=workspace["id"], pack_artifact_id=artifact["id"],
        target_url="https://x.test/apply", target_domain="x.test",
        ats_adapter_id="generic", ats_adapter_version="generic@1",
    )
    token = mint_session_token(conn, handoff_session_id=session["id"])

    resolved = resolve_session_scope(conn, raw_token=token)

    assert resolved == SessionScope(
        account_id="account_local",
        handoff_session_id=session["id"],
        workspace_id=workspace["id"],
        pack_artifact_id=artifact["id"],
    )
    conn.close()


def test_resolve_session_scope_refreshes_last_activity_at(tmp_path):
    conn = _conn(tmp_path)
    scope = _scope(tmp_path)
    workspace, artifact = _workspace_with_pack(conn)
    session = start_handoff_session(
        conn, scope, workspace_id=workspace["id"], pack_artifact_id=artifact["id"],
        target_url="https://x.test/apply", target_domain="x.test",
        ats_adapter_id="generic", ats_adapter_version="generic@1",
    )
    token = mint_session_token(conn, handoff_session_id=session["id"])
    original_activity = conn.execute(
        "SELECT last_activity_at FROM handoff_sessions WHERE id = ?", (session["id"],),
    ).fetchone()["last_activity_at"]

    # Force the stored activity timestamp into the past (but still within
    # the inactivity timeout) so a refresh is observably different, rather
    # than racing the clock within the same test process tick.
    stale_but_not_expired = (datetime.now(timezone.utc) - timedelta(minutes=30)).isoformat()
    conn.execute(
        "UPDATE handoff_sessions SET last_activity_at = ? WHERE id = ?",
        (stale_but_not_expired, session["id"]),
    )
    conn.commit()

    resolve_session_scope(conn, raw_token=token)

    refreshed_activity = conn.execute(
        "SELECT last_activity_at FROM handoff_sessions WHERE id = ?", (session["id"],),
    ).fetchone()["last_activity_at"]
    assert refreshed_activity != stale_but_not_expired
    assert refreshed_activity != original_activity
    conn.close()


def test_resolve_session_scope_rejects_unrecognized_token(tmp_path):
    conn = _conn(tmp_path)
    try:
        resolve_session_scope(conn, raw_token="not-a-real-token")
        assert False, "expected HandoffSessionTokenInvalid"
    except HandoffSessionTokenInvalid:
        pass
    conn.close()


def test_resolve_session_scope_rejects_expired_session(tmp_path):
    conn = _conn(tmp_path)
    scope = _scope(tmp_path)
    workspace, artifact = _workspace_with_pack(conn)
    session = start_handoff_session(
        conn, scope, workspace_id=workspace["id"], pack_artifact_id=artifact["id"],
        target_url="https://x.test/apply", target_domain="x.test",
        ats_adapter_id="generic", ats_adapter_version="generic@1",
    )
    token = mint_session_token(conn, handoff_session_id=session["id"])

    stale = (datetime.now(timezone.utc) - timedelta(hours=3)).isoformat()
    conn.execute(
        "UPDATE handoff_sessions SET last_activity_at = ? WHERE id = ?",
        (stale, session["id"]),
    )
    conn.commit()

    try:
        resolve_session_scope(conn, raw_token=token)
        assert False, "expected HandoffSessionExpired"
    except HandoffSessionExpired:
        pass
    conn.close()


def test_resume_handoff_session_succeeds_and_rotates_token(tmp_path):
    conn = _conn(tmp_path)
    scope = _scope(tmp_path)
    workspace, artifact = _workspace_with_pack(conn)
    session = start_handoff_session(
        conn, scope, workspace_id=workspace["id"], pack_artifact_id=artifact["id"],
        target_url="https://x.test/apply", target_domain="x.test",
        ats_adapter_id="generic", ats_adapter_version="generic@1",
    )
    original_token = mint_session_token(conn, handoff_session_id=session["id"])

    rotated_token = resume_handoff_session(conn, scope, handoff_session_id=session["id"])

    assert rotated_token != original_token
    # The original token must no longer resolve — resume revokes every
    # prior live token before minting the new one (design spec Section 3.2).
    try:
        resolve_session_scope(conn, raw_token=original_token)
        assert False, "expected the pre-resume token to be revoked"
    except HandoffSessionTokenInvalid:
        pass
    # The rotated token resolves successfully.
    resolved = resolve_session_scope(conn, raw_token=rotated_token)
    assert resolved.handoff_session_id == session["id"]
    conn.close()


def test_resume_handoff_session_rejects_expired_session(tmp_path):
    """The approved session-security design requires the server, not the
    browser extension, to be the authority on expiry: resume must
    independently reject a session whose last_activity_at exceeds the
    inactivity timeout, even when called directly (i.e. even if a client
    never pre-filtered candidates by staleness)."""
    conn = _conn(tmp_path)
    scope = _scope(tmp_path)
    workspace, artifact = _workspace_with_pack(conn)
    session = start_handoff_session(
        conn, scope, workspace_id=workspace["id"], pack_artifact_id=artifact["id"],
        target_url="https://x.test/apply", target_domain="x.test",
        ats_adapter_id="generic", ats_adapter_version="generic@1",
    )

    stale = (datetime.now(timezone.utc) - timedelta(hours=3)).isoformat()
    conn.execute(
        "UPDATE handoff_sessions SET last_activity_at = ? WHERE id = ?",
        (stale, session["id"]),
    )
    conn.commit()

    try:
        resume_handoff_session(conn, scope, handoff_session_id=session["id"])
        assert False, "expected HandoffSessionExpired"
    except HandoffSessionExpired:
        pass
    conn.close()


def test_discover_resumable_sessions_excludes_expired_session(tmp_path):
    """Discovery must not present a stale/expired session as resumable —
    the approved design places expiry authority on the server, and a
    session discovery offers as a resume candidate is exactly what a
    client-side orchestration will act on. An expired session lingering
    in discover results would defeat that authority even though resume
    itself would separately reject it."""
    conn = _conn(tmp_path)
    scope = _scope(tmp_path)
    workspace, artifact = _workspace_with_pack(conn)
    session = start_handoff_session(
        conn, scope, workspace_id=workspace["id"], pack_artifact_id=artifact["id"],
        target_url="https://x.test/apply", target_domain="x.test",
        ats_adapter_id="generic", ats_adapter_version="generic@1",
    )

    stale = (datetime.now(timezone.utc) - timedelta(hours=3)).isoformat()
    conn.execute(
        "UPDATE handoff_sessions SET last_activity_at = ? WHERE id = ?",
        (stale, session["id"]),
    )
    conn.commit()

    found = discover_resumable_handoff_sessions(
        conn, scope, workspace_id=workspace["id"], target_domain="x.test",
    )
    assert found == []
    conn.close()


def test_discover_resumable_sessions_includes_active_unexpired_session(tmp_path):
    """Companion to the exclusion test above: a session within the
    inactivity window must still surface normally, so the expiry filter
    cannot be so aggressive that it hides legitimately resumable
    sessions."""
    conn = _conn(tmp_path)
    scope = _scope(tmp_path)
    workspace, artifact = _workspace_with_pack(conn)
    session = start_handoff_session(
        conn, scope, workspace_id=workspace["id"], pack_artifact_id=artifact["id"],
        target_url="https://x.test/apply", target_domain="x.test",
        ats_adapter_id="generic", ats_adapter_version="generic@1",
    )

    stale_but_not_expired = (datetime.now(timezone.utc) - timedelta(minutes=30)).isoformat()
    conn.execute(
        "UPDATE handoff_sessions SET last_activity_at = ? WHERE id = ?",
        (stale_but_not_expired, session["id"]),
    )
    conn.commit()

    found = discover_resumable_handoff_sessions(
        conn, scope, workspace_id=workspace["id"], target_domain="x.test",
    )
    assert len(found) == 1
    assert found[0]["id"] == session["id"]
    conn.close()


def test_resolve_session_scope_rejects_non_in_progress_session(tmp_path):
    conn = _conn(tmp_path)
    scope = _scope(tmp_path)
    workspace, artifact = _workspace_with_pack(conn)
    session = start_handoff_session(
        conn, scope, workspace_id=workspace["id"], pack_artifact_id=artifact["id"],
        target_url="https://x.test/apply", target_domain="x.test",
        ats_adapter_id="generic", ats_adapter_version="generic@1",
    )
    token = mint_session_token(conn, handoff_session_id=session["id"])
    session_scope = resolve_session_scope(conn, raw_token=token)
    confirm_handoff_submission(conn, session_scope, handoff_session_id=session["id"])

    try:
        resolve_session_scope(conn, raw_token=token)
        assert False, "expected HandoffSessionTokenInvalid"
    except HandoffSessionTokenInvalid:
        pass
    conn.close()


def test_rotate_session_token_revokes_prior_tokens(tmp_path):
    conn = _conn(tmp_path)
    scope = _scope(tmp_path)
    workspace, artifact = _workspace_with_pack(conn)
    session = start_handoff_session(
        conn, scope, workspace_id=workspace["id"], pack_artifact_id=artifact["id"],
        target_url="https://x.test/apply", target_domain="x.test",
        ats_adapter_id="generic", ats_adapter_version="generic@1",
    )
    token_1 = mint_session_token(conn, handoff_session_id=session["id"])
    token_2 = rotate_session_token(conn, handoff_session_id=session["id"])

    assert token_1 != token_2
    try:
        resolve_session_scope(conn, raw_token=token_1)
        assert False, "expected HandoffSessionTokenInvalid for the revoked token"
    except HandoffSessionTokenInvalid:
        pass
    resolve_session_scope(conn, raw_token=token_2)  # succeeds
    conn.close()


def _started_session_scope(conn, scope, workspace, artifact, **overrides):
    body = {
        "target_url": "https://x.test/apply", "target_domain": "x.test",
        "ats_adapter_id": "generic", "ats_adapter_version": "generic@1",
    }
    body.update(overrides)
    session = start_handoff_session(
        conn, scope, workspace_id=workspace["id"], pack_artifact_id=artifact["id"], **body,
    )
    token = mint_session_token(conn, handoff_session_id=session["id"])
    return session, resolve_session_scope(conn, raw_token=token)


from webapp.services.handoff import HandoffDocumentKindUnsupported, fetch_session_document

_RENDERABLE_PACK_PAYLOAD = {
    "schema_version": "application-pack.v0",
    "job": {"company": "Acme", "title": "Engineer"},
    "cv_content": [
        {"unit_id": "cv_1", "unit_type": "cv_summary_line", "text": "Summary line."},
    ],
    "cover_letter_content": [
        {"unit_id": "cl_1", "unit_type": "cover_letter_paragraph", "text": "Paragraph."},
    ],
}


def _workspace_with_renderable_pack(conn, *, workspace_id="ws_docs", payload=None):
    ensure_profile_workspace(conn, account_id="account_local")
    workspace = create_workspace(
        conn, company="Acme", title="Engineer", account_id="account_local",
        workspace_id=workspace_id,
    )
    artifact = save_artifact(
        conn, workspace_id=workspace["id"], artifact_type="application_pack",
        payload=payload if payload is not None else _RENDERABLE_PACK_PAYLOAD,
    )
    return workspace, artifact


def test_fetch_session_document_renders_cv_from_the_pinned_pack(tmp_path):
    conn = _conn(tmp_path)
    scope = _scope(tmp_path)
    workspace, artifact = _workspace_with_renderable_pack(conn)
    _session, session_scope = _started_session_scope(conn, scope, workspace, artifact)

    rendered = fetch_session_document(
        conn, session_scope, kind="cv", documents_root=tmp_path / "documents",
    )
    assert rendered.kind == "cv"
    assert rendered.content
    assert rendered.filename.endswith(".docx")
    conn.close()


def test_fetch_session_document_renders_cover_letter_from_the_pinned_pack(tmp_path):
    conn = _conn(tmp_path)
    scope = _scope(tmp_path)
    workspace, artifact = _workspace_with_renderable_pack(conn)
    _session, session_scope = _started_session_scope(conn, scope, workspace, artifact)

    rendered = fetch_session_document(
        conn, session_scope, kind="cover_letter", documents_root=tmp_path / "documents",
    )
    assert rendered.kind == "cover_letter"
    assert rendered.content
    conn.close()


def test_fetch_session_document_rejects_unsupported_kind(tmp_path):
    conn = _conn(tmp_path)
    scope = _scope(tmp_path)
    workspace, artifact = _workspace_with_renderable_pack(conn)
    _session, session_scope = _started_session_scope(conn, scope, workspace, artifact)

    try:
        fetch_session_document(
            conn, session_scope, kind="resume_pdf", documents_root=tmp_path / "documents",
        )
        assert False, "expected HandoffDocumentKindUnsupported"
    except HandoffDocumentKindUnsupported:
        pass
    conn.close()


def test_fetch_session_document_ignores_a_newer_pack_confirmed_after_session_start(tmp_path):
    """The regression test for the exact invariant this endpoint exists
    to protect: a session started against pack A must keep rendering
    pack A's documents forever, even after a newer pack B becomes the
    workspace's current pack."""
    conn = _conn(tmp_path)
    scope = _scope(tmp_path)
    workspace, artifact_a = _workspace_with_renderable_pack(
        conn, payload={
            **_RENDERABLE_PACK_PAYLOAD,
            "cv_content": [
                {"unit_id": "cv_a", "unit_type": "cv_summary_line", "text": "Original pack summary."},
            ],
        },
    )
    _session, session_scope = _started_session_scope(conn, scope, workspace, artifact_a)

    # A newer pack for the SAME workspace, never referenced by this session.
    save_artifact(
        conn, workspace_id=workspace["id"], artifact_type="application_pack",
        payload={
            **_RENDERABLE_PACK_PAYLOAD,
            "cv_content": [
                {"unit_id": "cv_b", "unit_type": "cv_summary_line", "text": "Newer pack summary."},
            ],
        },
    )

    rendered = fetch_session_document(
        conn, session_scope, kind="cv", documents_root=tmp_path / "documents",
    )
    from docx import Document
    from io import BytesIO
    text = "\n".join(p.text for p in Document(BytesIO(rendered.content)).paragraphs)
    assert "Original pack summary" in text
    assert "Newer pack summary" not in text
    conn.close()


def test_fetch_session_document_cross_workspace_session_cannot_read_another_workspaces_pack(tmp_path):
    conn = _conn(tmp_path)
    scope = _scope(tmp_path)
    workspace_a, artifact_a = _workspace_with_renderable_pack(
        conn, workspace_id="ws_docs_a", payload={
            **_RENDERABLE_PACK_PAYLOAD,
            "cv_content": [
                {"unit_id": "cv_a", "unit_type": "cv_summary_line", "text": "Workspace A summary."},
            ],
        },
    )
    workspace_b, artifact_b = _workspace_with_renderable_pack(
        conn, workspace_id="ws_docs_b", payload={
            **_RENDERABLE_PACK_PAYLOAD,
            "cv_content": [
                {"unit_id": "cv_b", "unit_type": "cv_summary_line", "text": "Workspace B summary."},
            ],
        },
    )
    _session_a, scope_a = _started_session_scope(conn, scope, workspace_a, artifact_a)
    _session_b, scope_b = _started_session_scope(conn, scope, workspace_b, artifact_b)

    from docx import Document
    from io import BytesIO
    rendered_a = fetch_session_document(conn, scope_a, kind="cv", documents_root=tmp_path / "documents")
    rendered_b = fetch_session_document(conn, scope_b, kind="cv", documents_root=tmp_path / "documents")
    text_a = "\n".join(p.text for p in Document(BytesIO(rendered_a.content)).paragraphs)
    text_b = "\n".join(p.text for p in Document(BytesIO(rendered_b.content)).paragraphs)
    assert "Workspace A summary" in text_a
    assert "Workspace B summary" in text_b
    conn.close()
