"""Server-side CV text extraction (Bundle 7 spec §15.3): DOCX with python-docx,
PDF with pypdf. There is no OCR; a scanned PDF yields no text."""
from __future__ import annotations

import io

from product.application_document_contract import DOCX_MEDIA_TYPE, PDF_MEDIA_TYPE


def extract_text(content: bytes, media_type: str) -> str:
    if media_type == DOCX_MEDIA_TYPE:
        import docx
        document = docx.Document(io.BytesIO(content))
        parts = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(p for p in parts if p.strip())
    if media_type == PDF_MEDIA_TYPE:
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(content))
        return "\n".join((page.extract_text() or "") for page in reader.pages).strip()
    raise ValueError(f"unsupported media type {media_type}")
