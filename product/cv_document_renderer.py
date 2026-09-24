"""Deterministic downstream renderer for a confirmed ``cv-document-model.v2``.

This module never queries the Evidence Profile, calls the content or
statement planner, calls an LLM, or reads Application Intelligence, Job Fit,
or a database. Its only input is an already-built ``cv-document-model.v2``
payload; its only output is deterministic DOCX bytes. It renders exactly the
sections, ordering, and statement text the document model already decided,
and adds presentation-only structure (headings, bullets, spacing) on top.
"""

from __future__ import annotations

import zipfile
from io import BytesIO
from typing import Any

from docx import Document
from docx.shared import Pt

from product.cv_document_model import CV_DOCUMENT_MODEL_VERSION

CV_DOCUMENT_RENDERER_VERSION = "cv-document-renderer.v1"

_FROZEN_ZIP_TIMESTAMP = (1980, 1, 1, 0, 0, 0)

_NON_ROLE_SECTION_KIND = "items"
_ROLE_SECTION_KIND = "roles"


class CvDocumentRendererError(RuntimeError):
    pass


def render_cv_document(document_model: dict[str, Any]) -> bytes:
    """Render ``document_model`` (a ``cv-document-model.v2`` payload) to DOCX bytes.

    Rendering the same model twice, including across a real wall-clock gap,
    produces byte-identical output: see ``_freeze_docx_bytes``.
    """

    _validate_schema(document_model)

    document = Document()
    _set_base_style(document)

    for section in document_model.get("sections", []):
        _render_section(document, section)

    buffer = BytesIO()
    document.save(buffer)
    return _freeze_docx_bytes(buffer.getvalue())


def _validate_schema(document_model: Any) -> None:
    if not isinstance(document_model, dict):
        raise CvDocumentRendererError("cv document model must be a dict")

    schema_version = document_model.get("schema_version")
    if schema_version is None:
        raise CvDocumentRendererError("cv document model is missing schema_version")
    if schema_version != CV_DOCUMENT_MODEL_VERSION:
        raise CvDocumentRendererError(
            f"unsupported cv document model schema version {schema_version!r}; "
            f"expected {CV_DOCUMENT_MODEL_VERSION!r}"
        )

    sections = document_model.get("sections", [])
    if not isinstance(sections, list):
        raise CvDocumentRendererError("cv document model sections must be a list")
    for section in sections:
        if not isinstance(section, dict):
            raise CvDocumentRendererError("cv document model section must be a dict")
        if not isinstance(section.get("section_id"), str):
            raise CvDocumentRendererError("cv document model section is missing section_id")
        if not isinstance(section.get("title"), str):
            raise CvDocumentRendererError("cv document model section is missing title")


def _set_base_style(document: Document) -> None:
    style = document.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    for section in document.sections:
        section.top_margin = section.bottom_margin = Pt(54)
        section.left_margin = section.right_margin = Pt(54)


def _add_heading(document: Document, text: str) -> None:
    heading = document.add_heading(text, level=1)
    heading.paragraph_format.space_before = Pt(12)
    heading.paragraph_format.space_after = Pt(6)


def _render_section(document: Document, section: dict[str, Any]) -> None:
    if _ROLE_SECTION_KIND in section:
        _render_role_section(document, section)
        return
    _render_flat_section(document, section)


def _render_flat_section(document: Document, section: dict[str, Any]) -> None:
    items = section.get(_NON_ROLE_SECTION_KIND, [])
    visible_items = [item for item in items if item.get("statement_text") or item.get("text")]
    if not visible_items:
        return
    _add_heading(document, section["title"])
    for item in visible_items:
        text = item.get("statement_text", item.get("text"))
        document.add_paragraph(text, style="List Bullet")


def _render_role_section(document: Document, section: dict[str, Any]) -> None:
    roles = [role for role in section.get(_ROLE_SECTION_KIND, []) if role.get("statements")]
    if not roles:
        return
    _add_heading(document, section["title"])
    for role in roles:
        for statement in role.get("statements", []):
            text = statement.get("statement_text")
            if text:
                document.add_paragraph(text, style="List Bullet")


def _freeze_docx_bytes(content: bytes) -> bytes:
    """Rewrite a saved DOCX so its bytes depend only on document content.

    ``python-docx`` (via ``zipfile.ZipFile.writestr``) stamps every ZIP entry
    with the wall-clock time at save, so otherwise-identical documents saved
    a second apart produce different bytes. Rewriting every entry with a
    fixed timestamp removes that source of nondeterminism while leaving the
    OOXML content, entry order, and compression untouched.
    """

    source = zipfile.ZipFile(BytesIO(content))
    frozen_buffer = BytesIO()
    with zipfile.ZipFile(frozen_buffer, "w", zipfile.ZIP_DEFLATED) as frozen:
        for item in source.infolist():
            frozen_info = zipfile.ZipInfo(item.filename, date_time=_FROZEN_ZIP_TIMESTAMP)
            frozen_info.compress_type = item.compress_type
            frozen_info.external_attr = item.external_attr
            # ZIP creator metadata otherwise defaults to the host OS (0 on
            # Windows, 3 on Unix), making equivalent DOCX bytes differ across
            # platforms. Keep the frozen archive's metadata host-independent.
            frozen_info.create_system = 0
            frozen.writestr(frozen_info, source.read(item.filename))
    return frozen_buffer.getvalue()
